import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))

import export_i14y_json  # noqa: E402


class I14yRowExtractionTests(unittest.TestCase):
    def test_ifc_set_json_list_is_normalized_to_multiple_index_values(self):
        raw = '["Pset_PipeSegmentTypeCommon", "Qto_PipeSegmentBaseQuantities"]'
        self.assertEqual(
            ['Pset_PipeSegmentTypeCommon', 'Qto_PipeSegmentBaseQuantities'],
            export_i14y_json.parse_single_or_json_string_list(
                raw,
                'Properties',
                60,
                'IfcPropertySet (Pset) / IfcQuantitySet (Qto)',
            ),
        )

    def test_formula_only_template_row_is_not_exported_as_class(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Classes'
        sheet.cell(1, 1).value = 'Class-ID'
        sheet.cell(1, 2).value = 'Class-Code'
        sheet.cell(1, 3).value = 'Designation (EN)'
        sheet.cell(1, 4).value = 'Class-Assignment'
        sheet.cell(7, 4).value = '=IFERROR(INDEX(Rules!$B:$B,MATCH(B7,Rules!$A:$A,0)),"")'
        sheet.cell(8, 1).value = 'site'
        sheet.cell(8, 2).value = 'SITE'
        sheet.cell(8, 3).value = 'Site'

        rows = export_i14y_json.rowdict(sheet)

        self.assertEqual([row['_row'] for row in rows], [8])
        self.assertEqual(rows[0]['Class-ID'], 'site')


if __name__ == '__main__':
    unittest.main()
