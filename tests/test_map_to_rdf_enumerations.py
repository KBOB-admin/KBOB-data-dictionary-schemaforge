import tempfile
import unittest
from pathlib import Path
import shutil
import sys

from openpyxl import load_workbook
from rdflib import Graph, Literal, Namespace, URIRef
from rdflib.namespace import RDF, RDFS, SKOS, DCTERMS, OWL, DCAT, PROV

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))

import map_to_rdf  # noqa: E402


DD = Namespace('https://lindas.admin.ch/fobl/kbob/dd-fm/vocab/')
BSDD_DEF = Namespace('https://bsdd.buildingsmart.org/def#')
KBOB_BASE = 'https://lindas.admin.ch/fobl/kbob/dd-fm/0.3.0'
BASE = 'https://lindas.admin.ch/fobl/kbob/dd-fm/ifma/data_dictionary_area_management/0.5.0'


class EnumerationMappingTests(unittest.TestCase):
    def test_ifc_set_json_list_creates_separate_rdf_relations(self):
        graph = Graph()
        prop = URIRef('https://example.com/property/length')
        raw = '["Pset_PipeSegmentTypeCommon", "Qto_PipeSegmentBaseQuantities"]'

        map_to_rdf.add_ifc_set_links(graph, prop, raw, 60)

        linked = set(graph.objects(prop, DD.alignedWithIfcPset))
        self.assertEqual(
            {Literal('Pset_PipeSegmentTypeCommon'), Literal('Qto_PipeSegmentBaseQuantities')},
            linked,
        )

    def test_ifc_set_newline_list_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'must use a JSON list'):
            map_to_rdf.parse_single_or_json_string_list(
                'Pset_PipeSegmentTypeCommon\nQto_PipeSegmentBaseQuantities',
                'Properties',
                60,
                'IfcPropertySet (Pset) / IfcQuantitySet (Qto)',
            )

    def test_document_type_taxonomy_roles_and_nearest_parents(self):
        rows = [
            {'Class-Code': 'O', 'Beschreibung (DE)': 'Hierarchiestufe O', 'Category/Kategorie/Catégorie/Categoria': 'Document type taxonomy'},
            {'Class-Code': 'O02', 'Beschreibung (DE)': 'Hierarchiestufe O02', 'Category/Kategorie/Catégorie/Categoria': 'Document type taxonomy'},
            {'Class-Code': 'O02001', 'Beschreibung (DE)': 'Projektauftrag', 'Category/Kategorie/Catégorie/Categoria': 'Document type taxonomy'},
            {'Class-Code': 'O9', 'Beschreibung (DE)': 'Hierarchiestufe O9', 'Category/Kategorie/Catégorie/Categoria': 'Document type taxonomy'},
        ]
        roles, parents = map_to_rdf.class_taxonomy(rows)
        self.assertEqual(roles['O'], 'taxonomy-node')
        self.assertEqual(roles['O02'], 'taxonomy-node')
        self.assertEqual(roles['O02001'], 'document-type')
        self.assertEqual(roles['O9'], 'taxonomy-node')
        self.assertEqual(parents['O02'], 'O')
        self.assertEqual(parents['O02001'], 'O02')
        self.assertEqual(parents['O9'], 'O')

    def export_actual_workbooks(self):
        tmp = tempfile.TemporaryDirectory()
        out = Path(tmp.name)
        publication_inputs = out / 'publication-inputs'
        publication_inputs.mkdir()
        kbob_source = ROOT / 'templates' / 'Strukturvorlage_DataDictionary_KBOB_FM_v0.9.5.xlsx'
        area_source = ROOT / 'templates' / 'Strukturvorlage_AreaMgmt_v0.5.0.xlsx'
        kbob = publication_inputs / kbob_source.name
        area = publication_inputs / area_source.name
        shutil.copy2(kbob_source, kbob)
        shutil.copy2(area_source, area)
        approved_uris = {
            kbob: 'https://lindas.admin.ch/fobl/kbob/dd-fm/0.3.0/',
            area: 'https://lindas.admin.ch/fobl/kbob/dd-fm/ifma/data_dictionary_area_management/0.5.0/',
        }
        for workbook_path, approved_uri in approved_uris.items():
            workbook = load_workbook(workbook_path)
            header = workbook['Header']
            for row in range(1, header.max_row + 1):
                if header.cell(row, 1).value == 'DictionaryUri':
                    header.cell(row, 2).value = approved_uri
                    break
            else:
                workbook.close()
                raise AssertionError(f'DictionaryUri missing in {workbook_path}')
            workbook.save(workbook_path)
            workbook.close()
        map_to_rdf.main([
            str(kbob),
            str(area),
            '--out',
            str(out),
        ])
        graph = Graph()
        graph.parse(out / 'combined.ttl')
        return tmp, graph

    def test_resource_base_includes_dictionary_version(self):
        header = {
            'DictionaryUri': 'https://lindas.admin.ch/fobl/kbob/sbb/sbb_fdk_data_dictionary',
            'DictionaryVersion': '3.0.0',
        }
        self.assertEqual(
            map_to_rdf.resource_base_from_header(header),
            'https://lindas.admin.ch/fobl/kbob/sbb/sbb_fdk_data_dictionary/3.0.0',
        )
        already_versioned = {
            'DictionaryUri': 'https://lindas.admin.ch/fobl/kbob/sbb/sbb_fdk_data_dictionary/3.0.0',
            'DictionaryVersion': '3.0.0',
        }
        self.assertEqual(
            map_to_rdf.resource_base_from_header(already_versioned),
            'https://lindas.admin.ch/fobl/kbob/sbb/sbb_fdk_data_dictionary/3.0.0',
        )
        self.assertEqual(
            str(map_to_rdf.dictionary_version_uri_from_header(header)),
            'https://lindas.admin.ch/fobl/kbob/sbb/sbb_fdk_data_dictionary/3.0.0/',
        )

    def test_neutral_placeholder_is_preserved_until_publication_approval(self):
        header = {
            'DictionaryUri': 'https://example.com/sia/sia-416-2003',
            'DictionaryVersion': '0.0.1',
        }
        self.assertEqual(
            map_to_rdf.resource_base_from_header(header),
            'https://example.com/sia/sia-416-2003/0.0.1',
        )

    def test_approved_versioned_uri_is_preserved_without_double_version(self):
        header = {
            'DictionaryUri': 'https://lindas.admin.ch/fobl/kbob/dd-fm/ifma/data_dictionary_area_management/0.5.0/',
            'DictionaryVersion': '0.5.0',
        }
        self.assertEqual(
            map_to_rdf.resource_base_from_header(header),
            'https://lindas.admin.ch/fobl/kbob/dd-fm/ifma/data_dictionary_area_management/0.5.0',
        )
        self.assertEqual(
            str(map_to_rdf.dictionary_version_uri_from_header(header)),
            'https://lindas.admin.ch/fobl/kbob/dd-fm/ifma/data_dictionary_area_management/0.5.0/',
        )

    def test_missing_enumeration_id_uses_validator_canonical_suffix(self):
        schemes, _by_designation, diagnostics = map_to_rdf.build_enumerations([
            {
                '_row': 7,
                'Designation (EN)': 'Area type SIA 416 code',
                'Enumeration (EN)': '["1.2"]',
            }
        ])
        self.assertEqual(schemes[0]['id'], 'area-type-sia-416-code-enum')
        self.assertIn('area-type-sia-416-code-enum', diagnostics[0])

    def test_actual_workbooks_publish_schemes_and_values(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        scheme = URIRef(f'{BASE}/enumerations/area-type-sia-416-code-enum')
        value = URIRef(f'{BASE}/enumerations/area-type-sia-416-code-enum/values/0001-1-2')
        prop = URIRef(f'{BASE}/properties/sia-416-2003')

        self.assertIn((scheme, RDF.type, SKOS.ConceptScheme), graph)
        self.assertIn((scheme, RDF.type, DD.EnumerationScheme), graph)
        self.assertIn((value, RDF.type, DD.EnumerationValue), graph)
        self.assertIn((value, SKOS.inScheme, scheme), graph)
        self.assertIn((prop, DD.usesEnumerationScheme, scheme), graph)
        self.assertIn((prop, DD.hasAllowedValue, value), graph)
        self.assertIn((scheme, DD.status, Literal('Preview')), graph)
        self.assertNotIn((value, DD.status, Literal('Preview')), graph)
        self.assertEqual(len(list(graph.triples((None, SKOS.member, None)))), 0)

    def test_actual_workbooks_publish_data_templates_requirements_groups_and_documents(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        template = URIRef(f'{KBOB_BASE}/data-templates/grundstueck-dt')
        class_uri = URIRef(f'{KBOB_BASE}/classes/site')
        group_uri = URIRef(f'{KBOB_BASE}/group-of-properties/kbob-grundstuck-architekt')
        document_uri = URIRef(f'{KBOB_BASE}/documents/kbob-dk-fm-i14y-csv')
        req_guid = URIRef(f'{KBOB_BASE}/property-requirements/grundstueck-dt/kbob-grundstuck-architekt/guid')
        req_site_id = URIRef(f'{KBOB_BASE}/property-requirements/grundstueck-dt/kbob-grundstuck-architekt/site-id')
        prop_guid = URIRef(f'{KBOB_BASE}/properties/guid')
        dataset = URIRef(f'{KBOB_BASE}/')

        self.assertIn((template, RDF.type, DD.DataTemplate), graph)
        self.assertIn((dataset, DD.hasDataTemplate, template), graph)
        self.assertIn((template, DD.appliesToClass, class_uri), graph)
        self.assertNotIn((class_uri, DD.hasDataTemplate, template), graph)
        self.assertIn((template, DD.includesGroupOfProperties, group_uri), graph)
        self.assertIn((template, DD.referencesDocument, document_uri), graph)
        self.assertIn((document_uri, RDF.type, DD.Document), graph)
        self.assertIn((template, DD.hasPropertyRequirement, req_guid), graph)
        self.assertIn((req_guid, RDF.type, DD.PropertyRequirement), graph)
        self.assertIn((req_guid, DD.requirementOfDataTemplate, template), graph)
        self.assertIn((req_guid, DD.inGroupOfProperties, group_uri), graph)
        self.assertIn((req_guid, DD.requiresProperty, prop_guid), graph)
        self.assertIn((prop_guid, DD.requiredByTemplate, req_guid), graph)
        self.assertNotIn((template, DD.hasPropertyRequirement, req_site_id), graph)

    def test_dataset_uses_header_dictionary_name_and_source_workbook(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        dataset = URIRef(f'{KBOB_BASE}/')
        series = URIRef('https://lindas.admin.ch/fobl/kbob/dd-fm/')
        self.assertIn((dataset, RDFS.label, None), graph)
        self.assertIn((dataset, RDFS.label, Literal('KBOB FM Data Dictionary', lang='en')), graph)
        self.assertIn((dataset, DCTERMS.identifier, Literal('kbob_fm_data_dictionary')), graph)
        self.assertIn((dataset, OWL.versionInfo, Literal('0.3.0')), graph)
        source = URIRef(f'{KBOB_BASE}/source-workbooks/strukturvorlage-datadictionary-kbob-fm-v0-9-5-xlsx')
        self.assertIn((source, RDF.type, PROV.Entity), graph)
        self.assertIn((source, DCTERMS.identifier, Literal('Strukturvorlage_DataDictionary_KBOB_FM_v0.9.5.xlsx')), graph)
        self.assertIn((dataset, PROV.wasDerivedFrom, source), graph)
        self.assertIn((dataset, DCTERMS.source, source), graph)
        self.assertNotIn((dataset, DD.sourceWorkbook, None), graph)
        self.assertIn((dataset, RDF.type, DCAT.Dataset), graph)
        self.assertIn((dataset, RDF.type, DD.DataDictionary), graph)
        self.assertIn((series, RDF.type, URIRef('http://www.w3.org/ns/dcat#DatasetSeries')), graph)
        self.assertIn((dataset, DCTERMS.isVersionOf, series), graph)
        self.assertIn((dataset, URIRef('http://www.w3.org/ns/dcat#inSeries'), series), graph)
        self.assertIn((series, DCTERMS.hasVersion, dataset), graph)
        self.assertEqual(len(list(graph.subjects(RDF.type, DD.DataDictionary))), 2)
        self.assertFalse(any(str(subject).startswith('https://i14y.admin.ch/catalog/') for subject in graph.subjects(RDF.type, DD.DataDictionary)))

    def test_repeated_data_template_id_aggregates_zero_to_many_groups(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        template = URIRef(f'{KBOB_BASE}/data-templates/gebaeude-dt')
        groups = set(graph.objects(template, DD.includesGroupOfProperties))
        self.assertEqual(len(groups), 2)
        self.assertIn(
            URIRef(f'{KBOB_BASE}/group-of-properties/kbob-gebaude-architekt'),
            groups,
        )
        self.assertIn(
            URIRef(f'{KBOB_BASE}/group-of-properties/kbob-gebaude-betriebfm'),
            groups,
        )
        self.assertNotIn((None, DD.groupContainsProperty, None), graph)

        ifc_aligned_property = URIRef(f'{KBOB_BASE}/properties/above-ground')
        self.assertIn((ifc_aligned_property, BSDD_DEF.PropertySet, None), graph)
        self.assertNotIn((ifc_aligned_property, DD.inGroupOfProperties, None), graph)

        self.assertIn((URIRef(f'{KBOB_BASE}/classes/site'), DD.status, Literal('Candidate')), graph)
        self.assertIn((URIRef(f'{KBOB_BASE}/group-of-properties/kbob-grundstuck-architekt'), DD.status, Literal('Candidate')), graph)

    def test_property_datatype_ifc_datatype_and_status_are_published(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        prop = URIRef(f'{BASE}/properties/airconditioning')
        self.assertIn((prop, DD.datatype, Literal('BOOLEAN')), graph)
        self.assertIn((prop, DD.ifcDatatype, None), graph)
        self.assertIn((prop, DD.status, None), graph)

    def test_group_of_properties_uses_designation_as_label_and_description_as_definition(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        group = URIRef(f'{BASE}/group-of-properties/ifma-areamgmt')
        description_de = Literal(
            'Enthält alle flächenbezogenen Parameter gemäss Norm SIA 416. Sie definiert Hauptnutz-, Nebennutz- und Verkehrsflächen zur präzisen Steuerung beim Facility Management',
            lang='de',
        )

        self.assertIn((group, RDF.type, DD.GroupOfProperties), graph)
        self.assertIn((group, RDFS.label, Literal('IFMA_AreaMgmt', lang='en')), graph)
        self.assertIn((group, SKOS.prefLabel, Literal('IFMA_AreaMgmt', lang='en')), graph)
        self.assertNotIn((group, RDFS.label, description_de), graph)
        self.assertNotIn((group, SKOS.prefLabel, description_de), graph)
        self.assertIn((group, SKOS.definition, description_de), graph)
        self.assertIn((group, DCTERMS.description, description_de), graph)

    def test_x_requirement_retains_property_level_enumeration(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        requirement = URIRef(f'{BASE}/property-requirements/nettogeschossflaeche-dt/ifma-areamgmt/sia-2024-2021')
        scheme = URIRef(f'{BASE}/enumerations/room-usage-sia-2024-code-enum')
        self.assertIn((requirement, RDF.type, DD.PropertyRequirement), graph)
        self.assertIn((requirement, DD.usesEnumerationScheme, scheme), graph)
        self.assertIn((requirement, DD.usesCompleteValueScheme, None), graph)
        self.assertEqual(len(list(graph.objects(requirement, DD.allowedValue))), 0)

    def test_override_requirement_links_selected_existing_values(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        requirement = URIRef(f'{BASE}/property-requirements/umgebungsflaeche-dt/ifma-areamgmt/sia-416-2003')
        scheme = URIRef(f'{BASE}/enumerations/area-type-sia-416-code-enum')
        value = URIRef(f'{BASE}/enumerations/area-type-sia-416-code-enum/values/0001-1-2')
        self.assertIn((requirement, DD.usesEnumerationScheme, scheme), graph)
        self.assertIn((requirement, DD.allowedValue, value), graph)
        self.assertIn((value, RDF.type, DD.EnumerationValue), graph)

    def test_invalid_override_outside_linked_enumeration_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            bad_area = tmp / 'bad-area.xlsx'
            shutil.copyfile(ROOT / 'templates' / 'Strukturvorlage_AreaMgmt_v0.5.0.xlsx', bad_area)
            wb = load_workbook(bad_area)
            ws = wb['Data_Template']
            ws.cell(7, 8).value = '["NOT_A_VALUE"]'
            wb.save(bad_area)
            with self.assertRaisesRegex(ValueError, 'unresolved override value: NOT_A_VALUE'):
                map_to_rdf.main([
                    str(ROOT / 'templates' / 'Strukturvorlage_DataDictionary_KBOB_FM_v0.9.5.xlsx'),
                    str(bad_area),
                    '--out',
                    str(tmp / 'out'),
                ])

    def test_requirement_uris_are_deterministic_and_unique(self):
        tmp, graph = self.export_actual_workbooks()
        self.addCleanup(tmp.cleanup)

        requirements = list(graph.subjects(RDF.type, DD.PropertyRequirement))
        self.assertEqual(len(requirements), len(set(requirements)))
        self.assertIn(
            URIRef(f'{BASE}/property-requirements/nettogeschossflaeche-dt/ifma-areamgmt/sia-416-2003'),
            set(requirements),
        )

    def test_duplicate_lexical_values_are_distinguished_by_position(self):
        rows = [{
            'Enumeration-ID': 'room_designation_enum',
            'Enumeration-Code': 'ROOM',
            'Designation (EN)': 'Room designation',
            'Bezeichnung (DE)': 'Raumbezeichnung',
            'Enumeration (EN)': '["1", "2", "1"]',
            'Werteliste\n(DE)': '["HNF1", "HNF2", "NNF7"]',
            '_row': 7,
        }]
        schemes, by_designation, diagnostics = map_to_rdf.build_enumerations(rows)
        self.assertEqual(len(schemes[0]['values']), 3)
        self.assertIn('Room designation', by_designation)
        self.assertEqual(len(diagnostics), 1)

        uris = [
            map_to_rdf.mint_value_uri(BASE, schemes[0]['id'], value['position'], value['canonical'])
            for value in schemes[0]['values']
        ]
        self.assertEqual(len(set(uris)), 3)
        self.assertTrue(str(uris[0]).endswith('/values/0001-1'))
        self.assertTrue(str(uris[2]).endswith('/values/0003-1'))

    def test_misaligned_multilingual_value_lists_fail_fast(self):
        rows = [{
            'Enumeration-ID': 'bad_enum',
            'Designation (EN)': 'Bad enum',
            'Enumeration (EN)': '["A", "B"]',
            'Werteliste\n(DE)': '["A"]',
            '_row': 12,
        }]
        with self.assertRaisesRegex(ValueError, 'misaligned value-list lengths'):
            map_to_rdf.build_enumerations(rows)


if __name__ == '__main__':
    unittest.main()
