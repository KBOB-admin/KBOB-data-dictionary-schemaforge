#!/usr/bin/env python3
"""Export validated data dictionary workbooks to an I14Y-oriented JSON bundle.

This exporter does two things:
1. Emits API-aligned payload fragments for the I14Y Partner API based on the
   published OpenAPI spec captured from I14Y-ch/apiconsole.
2. Preserves semantic anchors that are critical for the PoC but are not first-
   class fields in the partner API payloads yet (IFC/QUDT/SKOS cross-links).

The resulting bundle is intentionally richer than a direct API POST body. The
`partnerApi` section contains strict request payload wrappers. The `semantic`
section preserves anchors and cross-links so they are never lost between
pipeline stages.
"""
from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_OUT = BASE_DIR / 'out' / 'i14y_bundle.json'
PARTNER_SPEC = BASE_DIR / 'external' / 'apiconsole__partner__v1__Release.json'
GLOBAL_CONFORMS_TO = [
    ('https://www.iso.org/standard/75401.html', 'ISO 23386'),
    ('https://www.iso.org/standard/85391.html', 'ISO 23387'),
    ('https://www.iso.org/standard/74932.html', 'ISO 12006-3'),
    ('https://www.dcat-ap.ch/releases/v3/3.0.0.html', 'DCAT-AP CH 3.0.0'),
    ('https://identifier.buildingsmart.org/uri/buildingsmart/ifc/4.3', 'IFC 4.3 schema'),
]


def display_workbook_path(path: Path) -> str:
    path = Path(path)
    try:
        return str(path.resolve().relative_to(BASE_DIR.resolve()))
    except Exception:
        return path.name


def slug(text: str) -> str:
    text = (text or '').strip().lower()
    repl = {
        'ä': 'ae', 'ö': 'oe', 'ü': 'ue', 'ß': 'ss',
        ' ': '-', '/': '-', '_': '-', '(': '', ')': '', ',': '-', '.': '-',
    }
    for a, b in repl.items():
        text = text.replace(a, b)
    text = re.sub(r'[^a-z0-9\-]+', '-', text)
    text = re.sub(r'-+', '-', text).strip('-')
    return text or 'item'


def version_segment(value: Any) -> str:
    text = str(value or '').strip()
    if re.fullmatch(r'\d+\.\d+\.\d+', text):
        return text
    return slug(text) if text else ''


def has_version_segment(uri: str) -> bool:
    tail = str(uri or '').rstrip('/').rsplit('/', 1)[-1]
    return bool(re.fullmatch(r'\d+\.\d+\.\d+', tail))


def parse_date(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    if re.match(r'^\d{4}-\d{2}-\d{2}$', s):
        return s + 'T00:00:00+00:00'
    return s


def ml(de=None, en=None, fr=None, it=None, rm=None):
    data = {'de': de, 'en': en, 'fr': fr, 'it': it, 'rm': rm}
    return {k: v for k, v in data.items() if v not in (None, '')}


def code_model(code: str | None):
    return {'code': code} if code else None


def identifier_model(identifier: str | None):
    return {'identifier': identifier} if identifier else None


def email_model(email: str | None):
    return {'email': email} if email else None


def vcard_model(email: str | None, *, de_name: str | None = None, en_name: str | None = None, note_de: str | None = None, note_en: str | None = None):
    if not email:
        return None
    obj = {'hasEmail': email}
    fn = ml(de=de_name, en=en_name)
    if fn:
        obj['fn'] = fn
    note = ml(de=note_de, en=note_en)
    if note:
        obj['note'] = note
    return obj


def resource_model(uri: str | None, de=None, en=None, fr=None, it=None):
    if not uri:
        return None
    obj = {'uri': uri}
    labels = ml(de=de, en=en, fr=fr, it=it)
    if labels:
        obj['label'] = labels
    return obj


def artifact_timestamp() -> str:
    return datetime.now(timezone.utc).strftime('%Y-%m-%dT%H-%M-%SZ')


def split_listish(value: Any) -> list[str]:
    if value in (None, ''):
        return []
    parts = re.split(r'[;,|]', str(value))
    return [p.strip() for p in parts if p and p.strip()]


def parse_single_or_json_string_list(value: Any, sheet_name: str, row_no: Any, column: str) -> list[str]:
    if value in (None, ''):
        return []
    if isinstance(value, (list, tuple)):
        parsed = list(value)
    else:
        text = str(value).strip()
        if not text:
            return []
        if text.startswith('['):
            try:
                parsed = json.loads(text)
            except json.JSONDecodeError as exc:
                raise ValueError(f'{sheet_name} row {row_no} column {column} must contain a valid JSON list') from exc
        else:
            if any(separator in text for separator in ('\n', '\r', ';')):
                raise ValueError(f'{sheet_name} row {row_no} column {column} must use a JSON list for multiple values')
            parsed = [text]
    if not isinstance(parsed, list) or not all(isinstance(item, str) and item.strip() for item in parsed):
        raise ValueError(f'{sheet_name} row {row_no} column {column} must contain one string or a JSON list of non-empty strings')
    normalized = [item.strip() for item in parsed]
    if len(set(normalized)) != len(normalized):
        raise ValueError(f'{sheet_name} row {row_no} column {column} contains duplicate references')
    return normalized


def first_nonempty(*vals):
    for v in vals:
        if v not in (None, ''):
            return v
    return None


def require_id(row: dict, column: str, sheet_name: str) -> str:
    value = str(row.get(column) or '').strip()
    if not value:
        row_no = row.get('_row', '?')
        raise ValueError(f'{sheet_name} row {row_no} is missing required system identifier {column}')
    return value


def owned_uri(base: str, subpath: str, identifier: str) -> str:
    identifier = str(identifier or '').strip()
    if not identifier:
        raise ValueError(f'Cannot mint {subpath} URI without an identifier')
    return f'{base}/{subpath}/{slug(identifier)}'


@dataclass
class WorkbookContext:
    path: Path
    wb: Any
    title: str
    publisher: str
    contact_email: str | None
    version: str
    release_date: str | None
    modified_date: str | None
    header: dict[str, Any]
    metadata: dict[str, Any]
    lindas_graph: str
    lindas_base: str
    dataset_identifier: str
    dataset_uri: str


@dataclass
class ExportOptions:
    contact_email: str | None = None
    landing_page_url: str | None = None
    workbook_download_url: str | None = None
    ttl_download_url: str | None = None
    json_download_url: str | None = None
    sparql_endpoint_url: str | None = None
    sparql_endpoint_description_url: str | None = None
    dataservice_landing_page_url: str | None = None
    include_distribution_media_types: bool = False


def read_dictionary_public(wb) -> dict[str, Any]:
    if 'Dictionary_public' not in wb.sheetnames:
        return {}
    ws = wb['Dictionary_public']
    meta = {}
    for row in range(4, ws.max_row + 1):
        field = ws.cell(row, 1).value
        value = ws.cell(row, 2).value
        if field and isinstance(field, str) and not field.startswith('──'):
            meta[field.strip()] = value
    return meta


def read_header(wb) -> dict[str, Any]:
    if 'Header' not in wb.sheetnames:
        return {}
    ws = wb['Header']
    meta = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] in (None, ''):
            continue
        key = str(row[0]).strip()
        value = row[1] if len(row) > 1 else None
        if value not in (None, ''):
            meta[key] = value
    return meta


def resource_base_from_header(header: dict[str, Any], publisher: str, lindas_graph: str, version: str | None = None) -> str:
    dictionary_uri = str(header.get('DictionaryUri') or '').strip().rstrip('/')
    if re.match(r'^https?://[^/]+(?:/|$)', dictionary_uri):
        family_base = dictionary_uri
    else:
        family_base = f'{lindas_graph}/bdch' if publisher == 'bdch' else lindas_graph
    segment = version_segment(version or header.get('DictionaryVersion'))
    if segment and not has_version_segment(family_base):
        return f'{family_base}/{segment}'
    return family_base


def infer_context(path: Path) -> WorkbookContext:
    wb = load_workbook(path, data_only=True)
    header = read_header(wb)
    meta = read_dictionary_public(wb)
    name = path.stem
    if 'kbob' in name.lower():
        title = 'KBOB FM Data Dictionary'
        publisher = 'kbob'
    elif 'area' in name.lower() or 'areamgmt' in name.lower():
        title = 'bDCH Area Management Data Dictionary'
        publisher = 'bdch'
    else:
        title = name
        publisher = slug(name)
    base_publisher = publisher
    title = first_nonempty(
        header.get('DictionaryName (EN)'),
        header.get('DictionaryName (DE)'),
        meta.get('Title'),
        meta.get('Name'),
        title,
    )
    publisher = first_nonempty(header.get('OrganizationCode'), meta.get('Owner / Publisher'), publisher)
    # Normalise known publisher tokens to canonical identifiers used by I14Y
    # e.g. 'KBOB' -> 'CH_KBOB'
    # Robustly normalise common publisher tokens to canonical identifiers used by I14Y
    if isinstance(publisher, str):
        pu = publisher.strip()
        up = pu.upper()
        if 'KBOB' in up or 'CH_KBOB' in up:
            publisher = 'CH_KBOB'
        elif 'BDCH' in up or 'CH_BDCH' in up:
            publisher = 'CH_BDCH'
        else:
            publisher = pu
    contact_email = meta.get('ContactEmail')
    version = first_nonempty(header.get('DictionaryVersion'), meta.get('Version'), '0.1.0')
    release_date = parse_date(meta.get('ReleaseDate'))
    modified_date = parse_date(meta.get('ModifiedDate')) or release_date
    lindas_graph = 'https://lindas.admin.ch/fobl/kbob/dd-fm'
    lindas_base = resource_base_from_header(header, base_publisher, lindas_graph, version)
    dataset_identifier = slug(first_nonempty(header.get('DictionaryCode'), name))
    dataset_uri = f'https://i14y.admin.ch/catalog/{dataset_identifier}'
    return WorkbookContext(path, wb, title, publisher, contact_email, version, release_date, modified_date, header, meta, lindas_graph, lindas_base, dataset_identifier, dataset_uri)


def override_contact_email(ctx: WorkbookContext, contact_email: str | None) -> WorkbookContext:
    if not contact_email:
        return ctx
    return WorkbookContext(
        path=ctx.path,
        wb=ctx.wb,
        title=ctx.title,
        publisher=ctx.publisher,
        contact_email=contact_email,
        version=ctx.version,
        release_date=ctx.release_date,
        modified_date=ctx.modified_date,
        header=ctx.header,
        metadata=ctx.metadata,
        lindas_graph=ctx.lindas_graph,
        lindas_base=ctx.lindas_base,
        dataset_identifier=ctx.dataset_identifier,
        dataset_uri=ctx.dataset_uri,
    )


def build_contact_point(ctx: WorkbookContext, contact_email: str | None):
    return vcard_model(
        contact_email,
        de_name=f'{ctx.title} Kontakt',
        en_name=f'{ctx.title} contact',
        note_de=f'Kontaktstelle fuer Metadaten und Publikation von {ctx.title}.',
        note_en=f'Contact point for metadata and publication of {ctx.title}.',
    )


def parse_language_codes(value: Any) -> list[str]:
    codes = []
    for item in split_listish(value):
        code = slug(item).replace('-', '')
        if code in {'de', 'en', 'fr', 'it', 'rm'} and code not in codes:
            codes.append(code)
    return codes


def infer_dataset_languages(ctx: WorkbookContext) -> list[str]:
    meta = ctx.metadata
    languages = []
    for code in parse_language_codes(meta.get('PrimaryLanguage')):
        if code not in languages:
            languages.append(code)
    for code in parse_language_codes(meta.get('MetadataLanguages')):
        if code not in languages:
            languages.append(code)
    for code, field in [('de', 'DictionaryName (DE)'), ('en', 'DictionaryName (EN)'), ('fr', 'DictionaryName (FR)'), ('it', 'DictionaryName (IT)'), ('rm', 'DictionaryName (RM)')]:
        if ctx.header.get(field) not in (None, '') and code not in languages:
            languages.append(code)
    return languages or ['en']


def build_keyword_models(ctx: WorkbookContext) -> list[dict[str, Any]]:
    meta = ctx.metadata
    labels: dict[str, dict[str, str]] = {}
    for lang, field in [('de', 'Keywords (DE)'), ('fr', 'Keywords (FR)'), ('en', 'Keywords (EN)')]:
        for raw in split_listish(meta.get(field)):
            key = raw.strip().lower()
            if not key:
                continue
            labels.setdefault(key, {})
            labels[key][lang] = raw.strip()
    result = []
    for key in sorted(labels):
        result.append({'label': labels[key]})
    return result


def generic_dataset_description(ctx: WorkbookContext) -> dict[str, str]:
    names = {
        'de': first_nonempty(ctx.header.get('DictionaryName (DE)'), ctx.header.get('DictionaryName (EN)'), ctx.title),
        'en': first_nonempty(ctx.header.get('DictionaryName (EN)'), ctx.header.get('DictionaryName (DE)'), ctx.title),
        'fr': first_nonempty(ctx.header.get('DictionaryName (FR)'), ctx.header.get('DictionaryName (EN)')),
        'it': first_nonempty(ctx.header.get('DictionaryName (IT)'), ctx.header.get('DictionaryName (EN)')),
    }
    return ml(
        de=f'Datenkatalog {names["de"]}.',
        en=f'Data dictionary {names["en"]}.',
        fr=f'Dictionnaire de données {names["fr"]}.' if names['fr'] else None,
        it=f'Dizionario dei dati {names["it"]}.' if names['it'] else None,
    )


def build_distribution(title: dict[str, str], description: dict[str, str], access_url: str, *, download_url: str | None = None, identifier: str | None = None, format_code: str | None = None, media_type_code: str | None = None, documentation_urls: list[str] | None = None, issued: str | None = None, modified: str | None = None, languages: list[str] | None = None) -> dict[str, Any]:
    payload = {
        'title': title,
        'description': description,
        'accessUrl': resource_model(access_url),
        'downloadUrl': resource_model(download_url) if download_url else None,
        'identifier': identifier,
        'format': code_model(format_code),
        'mediaType': code_model(media_type_code),
        'documentation': [resource_model(url) for url in (documentation_urls or []) if url],
        'issued': issued,
        'modified': modified,
        'languages': [code_model(code) for code in (languages or []) if code],
    }
    return {k: v for k, v in payload.items() if v not in (None, [], {})}


def build_dataset_distributions(ctx: WorkbookContext, options: ExportOptions) -> list[dict[str, Any]]:
    meta = ctx.metadata
    distributions = []
    languages = ['de', 'en']

    workbook_download_url = first_nonempty(
        options.workbook_download_url,
        meta.get('WorkbookDownloadURL'),
        meta.get('XLSX_Download_URL'),
        meta.get('SourceWorkbookUrl'),
    )
    if workbook_download_url:
        distributions.append(build_distribution(
            title=ml(de=f'{ctx.title} Excel-Strukturvorlage', en=f'{ctx.title} Excel workbook'),
            description=ml(
                de='Validierte Excel-Strukturvorlage des Data Dictionary.',
                en='Validated Excel workbook of the data dictionary.',
            ),
            access_url=workbook_download_url,
            download_url=workbook_download_url,
            identifier=f'{ctx.dataset_identifier}-xlsx',
            format_code='XLSX',
            media_type_code='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if options.include_distribution_media_types else None,
            issued=ctx.release_date,
            modified=ctx.modified_date,
            languages=languages,
        ))

    ttl_download_url = first_nonempty(options.ttl_download_url, meta.get('TTL_Download_URL'))
    if ttl_download_url:
        distributions.append(build_distribution(
            title=ml(de=f'{ctx.title} RDF/Turtle', en=f'{ctx.title} RDF/Turtle'),
            description=ml(
                de='RDF/Turtle-Export des publizierten Data Dictionary.',
                en='RDF/Turtle export of the published data dictionary.',
            ),
            access_url=ttl_download_url,
            download_url=ttl_download_url,
            identifier=f'{ctx.dataset_identifier}-ttl',
            format_code='TTL',
            media_type_code='text/turtle' if options.include_distribution_media_types else None,
            issued=ctx.release_date,
            modified=ctx.modified_date,
            languages=languages,
        ))

    json_download_url = first_nonempty(options.json_download_url, meta.get('JSON_Download_URL'))
    if json_download_url:
        distributions.append(build_distribution(
            title=ml(de=f'{ctx.title} JSON', en=f'{ctx.title} JSON'),
            description=ml(
                de='JSON-Export des publizierten Data Dictionary.',
                en='JSON export of the published data dictionary.',
            ),
            access_url=json_download_url,
            download_url=json_download_url,
            identifier=f'{ctx.dataset_identifier}-json',
            format_code='JSON',
            media_type_code='application/json' if options.include_distribution_media_types else None,
            issued=ctx.release_date,
            modified=ctx.modified_date,
            languages=languages,
        ))

    return distributions


def build_dataservice_payload(ctx: WorkbookContext, options: ExportOptions) -> dict[str, Any] | None:
    endpoint_url = first_nonempty(options.sparql_endpoint_url, ctx.metadata.get('SPARQL_Endpoint'))
    if not endpoint_url:
        return None

    contact_email = ctx.contact_email
    deputy_email = ctx.contact_email if ctx.contact_email and ctx.contact_email != contact_email else None
    endpoint_description_url = first_nonempty(options.sparql_endpoint_description_url, ctx.metadata.get('QualityAssuranceProcedureUrl'))
    landing_page_url = first_nonempty(options.dataservice_landing_page_url, options.landing_page_url, ctx.metadata.get('MoreInfoUrl'), ctx.dataset_uri)
    access_code = first_nonempty(ctx.metadata.get('AccessRights'), 'public')
    if isinstance(access_code, str):
        access_code = access_code.strip().upper()

    data = {
        'accessRights': code_model(access_code),
        'description': ml(
            de=f'SPARQL-Endpunkt fuer {ctx.title} auf dem Linked-Data-Dienst.',
            en=f'SPARQL endpoint for {ctx.title} on the linked data service.',
        ),
        'contactPoints': [cp for cp in [build_contact_point(ctx, contact_email)] if cp],
        'publisher': identifier_model(ctx.publisher),
        'title': ml(de=f'{ctx.title} SPARQL-Endpunkt', en=f'{ctx.title} SPARQL endpoint'),
        'identifiers': [f'{ctx.dataset_identifier}-sparql'],
        'issued': ctx.release_date,
        'modified': ctx.modified_date,
        'landingPages': [resource_model(landing_page_url, en='Data service landing page')],
        'endpointUrls': [resource_model(endpoint_url, en='SPARQL endpoint')],
        'endpointDescriptions': [resource_model(endpoint_description_url, en='Endpoint description')] if endpoint_description_url else None,
        'conformsTo': [x for x in [
            resource_model('https://www.w3.org/TR/sparql11-protocol/', en='SPARQL 1.1 Protocol'),
            resource_model(ctx.lindas_graph, en='Published LINDAS graph'),
        ] if x],
        'responsiblePerson': email_model(contact_email),
        'responsibleDeputy': email_model(deputy_email),
        'keywords': [
            {'label': ml(en='sparql')},
            {'label': ml(en='linked data')},
            {'label': ml(en='data service')},
        ],
        'version': ctx.version,
    }
    data = {k: v for k, v in data.items() if v not in (None, [], {})}
    return {'data': data}


def rowdict(ws, header_row=1, data_start=8):
    headers = {c: str(ws.cell(header_row, c).value).strip() for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value}
    rows = []

    def is_guidance_row(row_index: int) -> bool:
        vals = [ws.cell(row_index, c).value for c in range(1, ws.max_column + 1)]
        joined = ' | '.join(str(v).strip() for v in vals if v not in (None, '')).lower()
        markers = [
            'erforderlich für excel', 'erforderlich für merkmalkatalog', 'erforderlich für data dictionary',
            'sprachen', 'referenziert / system generiert', 'dropdown liste', 'validierung',
            'hinweise zur handhabung'
        ]
        return any(m in joined for m in markers)

    start = data_start
    for candidate in range(2, data_start + 1):
        vals = [ws.cell(candidate, c).value for c in range(1, ws.max_column + 1)]
        has_any = any(v not in (None, '') for v in vals)
        if has_any and not is_guidance_row(candidate):
            start = candidate
            break
    while start <= ws.max_row and is_guidance_row(start):
        start += 1

    semantic_anchors = {
        'Classes': ('Class-ID', 'Class-Code', 'Designation (EN)', 'Bezeichnung (DE)'),
        'Properties': ('Property-ID', 'Property-Code', 'Designation (EN)', 'Bezeichnung (DE)'),
        'GroupOfProperties': ('GoP-ID', 'GoP-Code', 'Designation (EN)', 'Bezeichnung (DE)'),
        'Values': ('Enumeration-ID', 'Designation (EN)', 'Bezeichnung (DE)', 'Enumeration (EN)'),
        'ConceptRelation': ('SourceConceptURI', 'SourceConceptID', 'TargetConceptURI', 'TargetConceptID'),
    }

    for r in range(start, ws.max_row + 1):
        rec = {}
        anyv = False
        for c, h in headers.items():
            v = ws.cell(r, c).value
            if v not in (None, ''):
                anyv = True
            rec[h] = v
        anchors = semantic_anchors.get(ws.title)
        has_semantic_content = any(
            rec.get(anchor) not in (None, '') for anchor in anchors
        ) if anchors else anyv
        if anyv and has_semantic_content:
            rec['_row'] = r
            rows.append(rec)
    return rows


def infer_concept_type(prop: dict, values_by_prop: dict[str, list[dict]]) -> tuple[str, dict[str, Any]]:
    code = str(prop.get('Property-Code') or '').strip()
    datatype = str(prop.get('DataType\n(Base Type)') or prop.get('DataType') or '').strip().upper()
    if values_by_prop.get(code):
        maxlen = max((len(str(v.get('Enumeration (EN)') or v.get('Bezeichnung (DE)') or '')) for v in values_by_prop[code]), default=0)
        return 'CodeList', {
            'codeListEntryDefaultSortProperty': 'Code',
            'codeListEntryValueType': 'String',
            'codeListEntryValueMaxLength': max(maxlen, 1),
        }
    if datatype in {'INTEGER', 'REAL', 'DECIMAL', 'NUMBER'} or 'MEASURE' in datatype:
        return 'Numeric', {}
    if datatype in {'BOOLEAN'}:
        return 'String', {'minLength': 1, 'maxLength': 5, 'pattern': '^(true|false|TRUE|FALSE|0|1)$'}
    maxlen = max(len(str(first_nonempty(prop.get('Designation (EN)'), prop.get('Bezeichnung (DE)'), code, 'x'))), 1)
    return 'String', {'minLength': 0, 'maxLength': maxlen * 4}


def build_dataset_payload(ctx: WorkbookContext, options: ExportOptions) -> dict[str, Any]:
    meta = ctx.metadata
    header = ctx.header
    descriptions = ml(
        de=meta.get('Description (DE)'),
        en=meta.get('Description (EN)'),
        fr=meta.get('Description (FR)'),
        it=meta.get('Description (IT)'),
    ) or generic_dataset_description(ctx)
    title_ml = ml(
        de=first_nonempty(meta.get('Title (DE)'), header.get('DictionaryName (DE)'), ctx.title),
        en=first_nonempty(meta.get('Title (EN)'), header.get('DictionaryName (EN)'), ctx.title),
        fr=first_nonempty(meta.get('Title (FR)'), header.get('DictionaryName (FR)')),
        it=first_nonempty(meta.get('Title (IT)'), header.get('DictionaryName (IT)')),
    )
    doc_urls = split_listish(meta.get('DocumentationURLs'))
    related_urls = split_listish(meta.get('RelatedResourceURLs'))

    # Construct safe identifiers list: include dataset identifier and optionally a
    # sanitized workbook stem (no paths or extensions). Avoid adding raw file paths
    # like 'templates/xxx.xlsx' which the I14Y API rejects.
    workbook_stem = ctx.path.stem if getattr(ctx, 'path', None) else ''
    workbook_identifier = slug(workbook_stem)
    identifiers = [ctx.dataset_identifier]
    if workbook_identifier and workbook_identifier != ctx.dataset_identifier:
        identifiers.append(workbook_identifier)

    # Normalize accessRights code to canonical uppercase token expected by the API
    access_code = first_nonempty(meta.get('AccessRights'), 'public')
    if isinstance(access_code, str):
        access_code = access_code.strip().upper()
    contact_email = ctx.contact_email
    deputy_email = ctx.contact_email if ctx.contact_email and ctx.contact_email != contact_email else None
    contact_points = [build_contact_point(ctx, contact_email)]
    landing_page_url = first_nonempty(options.landing_page_url, meta.get('MoreInfoUrl'))
    distributions = build_dataset_distributions(ctx, options)
    dataset_languages = infer_dataset_languages(ctx)
    keyword_models = build_keyword_models(ctx)
    data = {
        'accessRights': code_model(access_code),
        'description': descriptions,
        'contactPoints': [cp for cp in contact_points if cp],
        'publisher': identifier_model(ctx.publisher),
        'title': title_ml,
        # Identifiers are simple strings per DcatDatasetInputModel (array of string)
        'identifiers': identifiers,
        'issued': ctx.release_date,
        'modified': ctx.modified_date,
        'version': ctx.version,
        'versionNotes': ml(de=meta.get('VersionNotes (DE)'), en=meta.get('VersionNotes (EN)'), fr=meta.get('VersionNotes (FR)')),
        'landingPages': [resource_model(landing_page_url, en='Dataset landing page')] if landing_page_url else None,
        'documentation': [resource_model(u) for u in doc_urls if u],
        'relations': [resource_model(u) for u in related_urls if u],
        'distributions': distributions,
        'conformsTo': [x for x in (
            [resource_model(uri, en=label) for uri, label in GLOBAL_CONFORMS_TO] +
            [resource_model(meta.get('ConformsTo_DCAT_AP_CH'), en='DCAT-AP CH') if meta.get('ConformsTo_DCAT_AP_CH') else None]
        ) if x],
        'responsiblePerson': email_model(contact_email),
        'responsibleDeputy': email_model(deputy_email),
        'keywords': keyword_models,
        'languages': [code_model(code) for code in dataset_languages],
    }
    data = {k: v for k, v in data.items() if v not in (None, [], {})}
    # The partner API OpenAPI snapshot defines DcatDatasetInputModelDataWrapper
    # with exactly one lower-case `data` property and no extra properties.
    return {'data': data}


def build_concepts(ctx: WorkbookContext, classes: list[dict], props: list[dict], gops: list[dict], values: list[dict], concept_relations: list[dict]):
    by_prop_code: dict[str, list[dict]] = {str(p.get('Property-Code') or '').strip(): [] for p in props}
    for v in values:
        enum_bucket = str(v.get('Werteliste\n(DE)') or v.get('Liste de valeurs\n(FR)') or v.get('Lista valori\n(IT)') or v.get('Enumeration (EN)') or '').strip()
        for p in props:
            if str(p.get('EnumerationDesignation (EN)') or '').strip() == enum_bucket:
                by_prop_code[str(p.get('Property-Code') or '').strip()].append(v)
    concepts = []
    semantic_index = []
    for c in classes:
        class_id = require_id(c, 'Class-ID', 'Classes')
        owned = owned_uri(ctx.lindas_base, 'classes', class_id)
        payload = {
            'data': {
                'conceptType': 'String',
                'minLength': 0,
                'maxLength': 512,
                'name': ml(de=c.get('Bezeichnung (DE)'), en=c.get('Designation (EN)'), fr=c.get('Désignation (FR)'), it=c.get('Designazione (IT)')),
                'description': ml(de=c.get('Beschreibung (DE)'), en=c.get('Description (EN)'), fr=c.get('Description (FR)'), it=c.get('Descrizione (IT)')),
                'publisher': identifier_model(ctx.publisher),
                'responsiblePerson': email_model(ctx.contact_email),
                'version': str(c.get('Version date') or ctx.version),
                'identifiers': [class_id, str(owned)],
                'conformsTo': [x for x in [
                    resource_model(c.get('IFC_URI'), en='IFC 4.3 SSOT class pointer'),
                    resource_model(owned, en='LINDAS authoritative class IRI'),
                    resource_model(ctx.dataset_uri, en='I14Y dataset record')
                ] if x],
            }
        }
        concepts.append(payload)
        semantic_index.append({
            'kind': 'class',
            'id': class_id,
            'code': c.get('Class-Code'),
            'ownedUri': owned,
            'ifcUri': c.get('IFC_URI'),
            'ifcEntity': c.get('IfcObject Entity'),
            'ifcTypeObjectEntity': c.get('IfcTypeObject Entity'),
            'predefinedType': c.get('PredefinedType'),
            'sourceWorkbook': display_workbook_path(ctx.path),
            'i14yDatasetUri': ctx.dataset_uri,
        })
    for g in gops:
        gop_id = require_id(g, 'GoP-ID', 'GroupOfProperties')
        code = str(g.get('GoP-Code') or '').strip()
        owned = owned_uri(ctx.lindas_base, 'group-of-properties', gop_id)
        payload = {
            'data': {
                'conceptType': 'String',
                'minLength': 0,
                'maxLength': 512,
                'name': ml(de=g.get('Beschreibung (DE)'), en=g.get('Designation (EN)'), fr=g.get('Description (FR)'), it=g.get('Descrizione (IT)')),
                'description': ml(de=g.get('Beschreibung (DE)'), en=g.get('Description (EN)'), fr=g.get('Description (FR)'), it=g.get('Descrizione (IT)')),
                'publisher': identifier_model(ctx.publisher),
                'responsiblePerson': email_model(ctx.contact_email),
                'version': str(g.get('Version date') or ctx.version),
                'identifiers': [gop_id, owned],
                'conformsTo': [x for x in [resource_model(owned, en='LINDAS authoritative GoP IRI'), resource_model(ctx.dataset_uri, en='I14Y dataset record')] if x],
            }
        }
        concepts.append(payload)
        semantic_index.append({
            'kind': 'groupOfProperties',
            'id': gop_id,
            'code': code,
            'ownedUri': owned,
            'sourceWorkbook': display_workbook_path(ctx.path),
            'i14yDatasetUri': ctx.dataset_uri,
        })
    for p in props:
        prop_id = require_id(p, 'Property-ID', 'Properties')
        row_no = p.get('_row', '?')
        code = str(p.get('Property-Code') or '').strip()
        owned = owned_uri(ctx.lindas_base, 'properties', prop_id)
        ctype, extra = infer_concept_type(p, by_prop_code)
        payload_data = {
            'conceptType': ctype,
            'name': ml(de=p.get('Bezeichnung (DE)'), en=p.get('Designation (EN)'), fr=p.get('Désignation (FR)'), it=p.get('Designazione (IT)')),
            'description': ml(de=p.get('Beschreibung (DE)'), en=p.get('Description (EN)'), fr=p.get('Description (FR)'), it=p.get('Descrizione (IT)')),
            'publisher': identifier_model(ctx.publisher),
            'responsiblePerson': email_model(ctx.contact_email),
            'version': str(p.get('Version date') or ctx.version),
            'identifiers': [prop_id, str(owned)],
            'conformsTo': [x for x in [
                resource_model(p.get('IFC_URI'), en='IFC 4.3 SSOT property pointer'),
                resource_model(p.get('QUDT URI'), en='QUDT unit SSOT pointer'),
                resource_model(owned, en='LINDAS authoritative property IRI'),
                resource_model(ctx.dataset_uri, en='I14Y dataset record')
            ] if x],
        }
        payload_data.update(extra)
        concepts.append({'data': payload_data})
        semantic_index.append({
            'kind': 'property',
            'id': prop_id,
            'code': code,
            'ownedUri': owned,
            'ifcUri': p.get('IFC_URI'),
            'ifcPropertySet': parse_single_or_json_string_list(
                p.get('IfcPropertySet (Pset)\nIfcQuantitySet (Qto)'),
                'Properties',
                row_no,
                'IfcPropertySet (Pset) / IfcQuantitySet (Qto)',
            ),
            'qudtUri': p.get('QUDT URI'),
            'unitNameEn': p.get('Unit Name (EN)'),
            'dataType': p.get('DataType\n(Base Type)'),
            'dataTypeIfc': p.get('DataType\n(IFC)'),
            'sourceWorkbook': display_workbook_path(ctx.path),
            'i14yDatasetUri': ctx.dataset_uri,
            'enumerationDesignation': p.get('EnumerationDesignation (EN)'),
        })
    semantic_relations = []
    for r in concept_relations:
        semantic_relations.append({
            'sourcePropertyCode': r.get('PropertyCode'),
            'relationType': r.get('RelationType'),
            'relatedPropertyUri': r.get('RelatedPropertyUri'),
            'notes': r.get('Notes'),
            'sourceWorkbook': display_workbook_path(ctx.path),
        })
    codelist_entries = []
    for pcode, vals in by_prop_code.items():
        if not vals:
            continue
        entries = []
        for v in vals:
            value_id = require_id(v, 'Enumeration-ID', 'Values')
            entries.append({
                'code': value_id,
                'label': ml(de=v.get('Bezeichnung (DE)'), en=v.get('Enumeration (EN)'), fr=v.get('Désignation (FR)'), it=v.get('Designazione (IT)')),
                'ownedUri': owned_uri(ctx.lindas_base, 'values', value_id),
                'unit': v.get('Units/Einheit/Unité/Unità'),
            })
        codelist_entries.append({'propertyCode': pcode, 'entries': entries})
    return concepts, semantic_index, semantic_relations, codelist_entries


def extract_rows(ctx: WorkbookContext):
    wb = ctx.wb
    classes = rowdict(wb['Classes']) if 'Classes' in wb.sheetnames else []
    props = rowdict(wb['Properties']) if 'Properties' in wb.sheetnames else []
    gops = rowdict(wb['GroupOfProperties']) if 'GroupOfProperties' in wb.sheetnames else []
    values = rowdict(wb['Values']) if 'Values' in wb.sheetnames else []
    concept_relations = rowdict(wb['ConceptRelation'], header_row=2, data_start=5) if 'ConceptRelation' in wb.sheetnames else []
    return classes, props, gops, values, concept_relations


def build_bundle(workbooks: list[Path], options: ExportOptions | None = None):
    options = options or ExportOptions()
    partner_spec = json.loads(PARTNER_SPEC.read_text()) if PARTNER_SPEC.exists() else None
    bundle = {
        'bundleVersion': '0.1.0',
        'generator': 'schemaforge/export_i14y_json.py',
        'partnerApiSpec': {
            'title': partner_spec['info']['title'],
            'version': partner_spec['info']['version'],
            'serverAbn': partner_spec['servers'][1]['url'] if partner_spec and len(partner_spec['servers']) > 1 else None,
            'serverProd': partner_spec['servers'][0]['url'] if partner_spec else None,
        } if partner_spec else None,
        'datasets': [],
        'dataservices': [],
        'concepts': [],
        'codeListEntries': [],
        'semantic': {'anchors': [], 'skosRelations': [], 'crossLinks': []}
    }
    for wb_path in workbooks:
        ctx = override_contact_email(infer_context(wb_path), options.contact_email)
        classes, props, gops, values, concept_relations = extract_rows(ctx)
        dataset_payload = build_dataset_payload(ctx, options)
        dataservice_payload = build_dataservice_payload(ctx, options)
        concepts, anchors, skos_relations, code_entries = build_concepts(ctx, classes, props, gops, values, concept_relations)
        bundle['datasets'].append({
            'sourceWorkbook': display_workbook_path(wb_path),
            'i14yRecordUri': ctx.dataset_uri,
            'lindasGraphUri': ctx.lindas_graph,
            'lindasResourceBaseUri': ctx.lindas_base,
            'partnerApi': dataset_payload,
            'counts': {
                'classes': len(classes),
                'properties': len(props),
                'groupOfProperties': len(gops),
                'values': len(values),
                'conceptRelations': len(concept_relations),
            },
        })
        if dataservice_payload:
            bundle['dataservices'].append({
                'sourceWorkbook': display_workbook_path(wb_path),
                'datasetIdentifier': ctx.dataset_identifier,
                'partnerApi': dataservice_payload,
            })
        bundle['concepts'].extend(concepts)
        bundle['codeListEntries'].extend(code_entries)
        bundle['semantic']['anchors'].extend(anchors)
        bundle['semantic']['skosRelations'].extend(skos_relations)
        bundle['semantic']['crossLinks'].append({
            'datasetUri': ctx.dataset_uri,
            'lindasGraphUri': ctx.lindas_graph,
            'lindasResourceBaseUri': ctx.lindas_base,
            'sourceWorkbook': display_workbook_path(wb_path),
            'note': 'I14Y record should reference LINDAS authoritative IRIs; LINDAS RDF should reference this I14Y dataset URI as provenance.'
        })
    return bundle


def validate_bundle(bundle: dict[str, Any]) -> list[str]:
    errs = []
    for ds in bundle.get('datasets', []):
        data = ds['partnerApi']['data']
        for req in ['accessRights', 'description', 'publisher', 'title']:
            if req not in data:
                errs.append(f'dataset missing required field: {req}')
    for concept in bundle.get('concepts', []):
        data = concept['data']
        for req in ['description', 'name', 'publisher', 'responsiblePerson', 'version']:
            if req not in data:
                errs.append(f'concept missing required field: {req}')
        for ct in data.get('conformsTo', []):
            if 'uri' in ct and ct['uri'] and not str(ct['uri']).startswith(('http://', 'https://')):
                errs.append(f'non-absolute URI in conformsTo: {ct["uri"]}')
    for a in bundle.get('semantic', {}).get('anchors', []):
        for k in ('ifcUri', 'qudtUri', 'ownedUri'):
            v = a.get(k)
            if v and not str(v).startswith(('http://', 'https://')):
                errs.append(f'anchor {k} is not absolute: {v}')
    return errs


def main(argv=None):
    ap = argparse.ArgumentParser(
        description='Export validated SchemaForge workbook(s) to an I14Y-oriented JSON bundle.'
    )
    ap.add_argument('workbooks', nargs='+', help='Path(s) to Strukturvorlage workbook file(s) to process. Provide one or more workbook file paths.')
    ap.add_argument('--out', default=str(DEFAULT_OUT), help='Output path for the generated JSON bundle.')
    ap.add_argument('--contact-email', default=None, help='Override contact email with an email address that already exists as a person in i14y.')
    ap.add_argument('--landing-page-url', default=None, help='Override the dataset landing page URL.')
    ap.add_argument('--workbook-download-url', default=None, help='Override the workbook download URL for a dataset distribution.')
    ap.add_argument('--ttl-download-url', default=None, help='Override the RDF/Turtle download URL for a dataset distribution.')
    ap.add_argument('--json-download-url', default=None, help='Override the JSON download URL for a dataset distribution.')
    ap.add_argument('--sparql-endpoint-url', default=None, help='Override the SPARQL endpoint URL for a linked data service export.')
    ap.add_argument('--sparql-endpoint-description-url', default=None, help='Override the SPARQL endpoint description/documentation URL.')
    ap.add_argument('--dataservice-landing-page-url', default=None, help='Override the landing page URL used for the generated data service metadata.')
    args = ap.parse_args(argv)

    # Validate input files exist
    paths = [Path(p) for p in args.workbooks]
    missing = [str(p) for p in paths if not p.exists()]
    if missing:
        ap.error('The following workbook files were not found: ' + ', '.join(missing))

    options = ExportOptions(
        contact_email=args.contact_email,
        landing_page_url=args.landing_page_url,
        workbook_download_url=args.workbook_download_url,
        ttl_download_url=args.ttl_download_url,
        json_download_url=args.json_download_url,
        sparql_endpoint_url=args.sparql_endpoint_url,
        sparql_endpoint_description_url=args.sparql_endpoint_description_url,
        dataservice_landing_page_url=args.dataservice_landing_page_url,
    )
    bundle = build_bundle(paths, options=options)
    errs = validate_bundle(bundle)
    if errs:
        bundle['validationErrors'] = errs
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding='utf8')
    if bundle.get('datasets'):
        dataset_payload = bundle['datasets'][0]['partnerApi']
        dataset_ui_path = out.with_name('i14y_dataset_import_for_ui.json')
        dataset_ui_path.write_text(json.dumps(dataset_payload, ensure_ascii=False, indent=2), encoding='utf8')

        dataset_slug = slug(bundle['datasets'][0].get('i14yRecordUri', '').rsplit('/', 1)[-1] or 'dataset')
        stamped_name = f'i14y_dataset_import_for_ui__{dataset_slug}__{artifact_timestamp()}.json'
        stamped_path = out.with_name(stamped_name)
        stamped_path.write_text(json.dumps(dataset_payload, ensure_ascii=False, indent=2), encoding='utf8')

        manifest_path = out.with_name('i14y_dataset_import_for_ui__LATEST.txt')
        manifest_lines = [
            'Latest i14y dataset upload artifact',
            f'latest_file={dataset_ui_path.name}',
            f'versioned_file={stamped_path.name}',
            f'source_bundle={out.name}',
            f'generated_utc={artifact_timestamp()}',
        ]
        manifest_path.write_text('\n'.join(manifest_lines) + '\n', encoding='utf8')
    if bundle.get('dataservices'):
        dataservice_payload = bundle['dataservices'][0]['partnerApi']
        dataservice_path = out.with_name('i14y_dataservice_import_for_ui.json')
        dataservice_path.write_text(json.dumps(dataservice_payload, ensure_ascii=False, indent=2), encoding='utf8')

        service_slug = slug(bundle['dataservices'][0].get('datasetIdentifier') or 'dataservice')
        stamped_service_name = f'i14y_dataservice_import_for_ui__{service_slug}__{artifact_timestamp()}.json'
        stamped_service_path = out.with_name(stamped_service_name)
        stamped_service_path.write_text(json.dumps(dataservice_payload, ensure_ascii=False, indent=2), encoding='utf8')

        dataservice_manifest_path = out.with_name('i14y_dataservice_import_for_ui__LATEST.txt')
        dataservice_manifest_lines = [
            'Latest i14y data service upload artifact',
            f'latest_file={dataservice_path.name}',
            f'versioned_file={stamped_service_path.name}',
            f'source_bundle={out.name}',
            f'generated_utc={artifact_timestamp()}',
        ]
        dataservice_manifest_path.write_text('\n'.join(dataservice_manifest_lines) + '\n', encoding='utf8')
    print(out)
    print('validation_ok' if not errs else f'validation_errors {len(errs)}')
    if errs:
        for e in errs[:20]:
            print('-', e)


if __name__ == '__main__':
    import sys
    raise SystemExit(main(sys.argv[1:]))
