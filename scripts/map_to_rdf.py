#!/usr/bin/env python3
from rdflib import Graph, Namespace, URIRef, Literal
from rdflib.namespace import RDF, RDFS, SKOS, DCTERMS, OWL, XSD
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from pathlib import Path
import argparse, ast, json, re

LINDAS_BASE = 'https://lindas.admin.ch/fobl/kbob/dd-fm'
BDCH = 'bdch'
IFC_CLASS_BASE = 'https://identifier.buildingsmart.org/uri/buildingsmart/ifc/4.3/class'
QUDT = Namespace('http://qudt.org/schema/qudt/')
PROV = Namespace('http://www.w3.org/ns/prov#')
DD = Namespace('https://lindas.admin.ch/fobl/kbob/dd-fm/vocab/')
SCHEMA = Namespace('https://schema.org/')


def slug(s):
    s = str(s or '').strip().lower()
    repl = {' ':'-','ä':'ae','ö':'oe','ü':'ue','ß':'ss','/':'-','_':'-','(':'',')':''}
    for a,b in repl.items():
        s = s.replace(a,b)
    s = re.sub(r'[^a-z0-9\-]+','-',s)
    s = re.sub(r'-+','-',s).strip('-')
    return s or 'item'


def rowdict(ws, header_row=1, data_start=8):
    headers={str(ws.cell(header_row,c).value).strip():c for c in range(1,ws.max_column+1) if ws.cell(header_row,c).value}
    rows=[]

    def is_guidance_row(row_index: int) -> bool:
        vals = [ws.cell(row_index, c).value for c in range(1, ws.max_column + 1)]
        joined = ' | '.join(str(v).strip() for v in vals if v not in (None, '')).lower()
        markers = [
            'erforderlich für excel', 'erforderlich für merkmalkatalog', 'erforderlich für data dictionary',
            'sprachen', 'referenziert / system generiert', 'dropdown liste', 'validierung',
            'hinweise zur handhabung'
        ]
        return any(m in joined for m in markers)

    start = data_start
    for candidate in range(2, data_start + 1):
        vals = [ws.cell(candidate, c).value for c in range(1, ws.max_column + 1)]
        has_any = any(v not in (None, '') for v in vals)
        if has_any and not is_guidance_row(candidate):
            start = candidate
            break
    while start <= ws.max_row and is_guidance_row(start):
        start += 1

    for r in range(start, ws.max_row+1):
        rec={}
        anyv=False
        for h,c in headers.items():
            v=ws.cell(r,c).value
            if v not in (None,''):
                anyv=True
            rec[h]=v
        if anyv:
            rec['_row']=r
            rows.append(rec)
    return rows


def norm(value):
    if value in (None, ''):
        return ''
    return re.sub(r'\s+', ' ', str(value).strip()).casefold()


def require_id(row, column, sheet_name):
    value = str(row.get(column) or '').strip()
    if not value:
        row_no = row.get('_row', '?')
        raise ValueError(f'{sheet_name} row {row_no} is missing required system identifier {column}')
    return value


def mint_from_id(base, subpath, identifier):
    identifier = str(identifier or '').strip()
    if not identifier:
        raise ValueError(f'Cannot mint {subpath} URI without an identifier')
    return URIRef(f"{base}/{subpath}/{slug(identifier)}")


def mint_value_uri(base, enum_id, position, lexical):
    return URIRef(f"{base}/enumerations/{slug(enum_id)}/values/{position:04d}-{slug(lexical)}")


def mint_requirement_uri(base, template_id, property_id):
    return URIRef(f"{base}/property-requirements/{slug(template_id)}/{slug(property_id)}")


def parse_value_list(value, sheet_name, row_no, column):
    if value in (None, ''):
        return []
    if isinstance(value, (list, tuple)):
        return list(value)
    text = str(value).strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        try:
            parsed = ast.literal_eval(text)
        except (ValueError, SyntaxError) as exc:
            raise ValueError(f'{sheet_name} row {row_no} column {column} must contain a JSON-style list') from exc
    if not isinstance(parsed, list):
        raise ValueError(f'{sheet_name} row {row_no} column {column} must contain a list')
    return parsed


def parse_override_list(value, sheet_name, row_no, column, diagnostics, path):
    try:
        return parse_value_list(value, sheet_name, row_no, column)
    except ValueError:
        text = str(value or '').strip()
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            parsed = text.strip('"').strip("'")
        if parsed in (None, '') or isinstance(parsed, list):
            raise
        diagnostics.append(
            f'{Path(path).name}: WARNING Data_Template row {row_no} column {column} uses scalar override '
            f'{text}; normalized to one-item list for mapping.'
        )
        return [parsed]


def lexical(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    return str(value).strip()


def cell_value(wb_data, wb_formula, sheet_name, row, col):
    data_value = wb_data[sheet_name].cell(row, col).value
    if data_value not in (None, ''):
        return data_value
    formula_value = wb_formula[sheet_name].cell(row, col).value
    if hasattr(formula_value, 'text') and hasattr(formula_value, 'ref'):
        text = str(formula_value.text)
        match = re.fullmatch(r'=TRANSPOSE\(([A-Za-z0-9_ ]+)!([A-Z]+[0-9]+):([A-Z]+[0-9]+)\)', text.strip())
        if match:
            ref_sheet, start_cell, end_cell = match.groups()
            min_col, min_row, max_col, max_row = range_boundaries(formula_value.ref)
            if min_row <= row <= max_row and min_col <= col <= max_col and ref_sheet in wb_formula.sheetnames:
                source_start_row, source_start_col = coordinate_to_tuple(start_cell)
                source_end_row, source_end_col = coordinate_to_tuple(end_cell)
                if source_start_col == source_end_col:
                    return wb_formula[ref_sheet].cell(source_start_row + (col - min_col), source_start_col).value
    if isinstance(formula_value, str) and formula_value.startswith('='):
        match = re.fullmatch(r'=([A-Za-z0-9_ ]+)!([A-Z]+[0-9]+)', formula_value.strip())
        if match and match.group(1) in wb_formula.sheetnames:
            ref_sheet, ref_cell = match.groups()
            ref_row, ref_col = coordinate_to_tuple(ref_cell)
            return wb_formula[ref_sheet].cell(ref_row, ref_col).value
    return formula_value


def typed_literal(value):
    if isinstance(value, bool):
        return Literal(value)
    if isinstance(value, int) and not isinstance(value, bool):
        return Literal(value, datatype=XSD.integer)
    if isinstance(value, float):
        return Literal(value, datatype=XSD.decimal)
    return None


def add_pref_labels(g, uri, de=None, en=None, fr=None, it=None):
    labels = [('en', en), ('de', de), ('fr', fr), ('it', it)]
    for lang, value in labels:
        if value not in (None, ''):
            lit = Literal(lexical(value), lang=lang)
            g.add((uri, SKOS.prefLabel, lit))
            g.add((uri, RDFS.label, lit))


def build_enumerations(values_rows, sheet_name='Values'):
    schemes = []
    by_designation = {}
    by_id = {}
    diagnostics = []
    for row in values_rows:
        enum_id = require_id(row, 'Enumeration-ID', sheet_name)
        row_no = row.get('_row', '?')
        designation_en = str(row.get('Designation (EN)') or '').strip()
        if not designation_en:
            raise ValueError(f'{sheet_name} row {row_no} is missing required Designation (EN)')
        if enum_id in by_id:
            raise ValueError(f'{sheet_name} row {row_no} duplicates Enumeration-ID: {enum_id}')
        lists = {
            'en': parse_value_list(row.get('Enumeration (EN)'), sheet_name, row_no, 'Enumeration (EN)'),
            'de': parse_value_list(row.get('Werteliste\n(DE)'), sheet_name, row_no, 'Werteliste (DE)'),
            'fr': parse_value_list(row.get('Liste de valeurs\n(FR)'), sheet_name, row_no, 'Liste de valeurs (FR)'),
            'it': parse_value_list(row.get('Lista valori\n(IT)'), sheet_name, row_no, 'Lista valori (IT)'),
        }
        populated_lengths = {lang: len(items) for lang, items in lists.items() if items}
        if not populated_lengths:
            raise ValueError(f'{sheet_name} row {row_no} has no enumeration values')
        canonical_len = populated_lengths.get('en') or next(iter(populated_lengths.values()))
        for lang, length in populated_lengths.items():
            if length != canonical_len:
                raise ValueError(
                    f'{sheet_name} row {row_no} has misaligned value-list lengths: '
                    f'{", ".join(f"{k}={v}" for k, v in sorted(populated_lengths.items()))}'
                )
        values = []
        seen_value_uris = set()
        for index in range(canonical_len):
            labels = {lang: (items[index] if items else None) for lang, items in lists.items()}
            canonical = lexical(labels.get('en') or labels.get('de') or labels.get('fr') or labels.get('it') or index + 1)
            value_uri_key = f'{index + 1:04d}-{slug(canonical)}'
            if value_uri_key in seen_value_uris:
                raise ValueError(f'{sheet_name} row {row_no} produces a duplicate value URI token: {value_uri_key}')
            seen_value_uris.add(value_uri_key)
            values.append({'position': index + 1, 'canonical': canonical, 'labels': labels, 'raw': labels.get('en')})
        scheme = {
            'id': enum_id,
            'code': str(row.get('Enumeration-Code') or '').strip(),
            'status': str(row.get('Status') or '').strip(),
            'designation_en': designation_en,
            'labels': {
                'en': row.get('Designation (EN)'),
                'de': row.get('Bezeichnung (DE)'),
                'fr': row.get('Désignation (FR)'),
                'it': row.get('Designazione (IT)'),
            },
            'values': values,
            'row': row,
        }
        by_id[enum_id] = scheme
        if designation_en in by_designation:
            raise ValueError(f'{sheet_name} row {row_no} duplicates Designation (EN): {designation_en}')
        by_designation[designation_en] = scheme
        diagnostics.append(f'{designation_en}: {len(values)} value(s), URI basis Enumeration-ID={enum_id}')
        schemes.append(scheme)
    return schemes, by_designation, diagnostics


def add_lookup(lookup, keys, value):
    for key in keys:
        n = norm(key)
        if n:
            lookup[n] = value


def build_matrix_model(wb_data, wb_formula, path, base, dataset_uri, classes, props, gops, docs, property_enum_schemes, enum_value_lookup):
    sheet_name = 'Data_Template'
    if sheet_name not in wb_data.sheetnames:
        return [], []
    ws = wb_data[sheet_name]
    diagnostics = []

    class_lookup = {}
    for row in classes:
        class_id = require_id(row, 'Class-ID', 'Classes')
        uri = mint_from_id(base, 'classes', class_id)
        add_lookup(class_lookup, [
            class_id, row.get('Class-Code'), row.get('Designation (EN)'), row.get('Bezeichnung (DE)'),
            row.get('Désignation (FR)'), row.get('Designazione (IT)')
        ], {'row': row, 'id': class_id, 'uri': uri})

    prop_lookup = {}
    for row in props:
        prop_id = require_id(row, 'Property-ID', 'Properties')
        uri = mint_from_id(base, 'properties', prop_id)
        add_lookup(prop_lookup, [
            prop_id, row.get('Property-Code'), row.get('Designation (EN)'), row.get('Bezeichnung (DE)'),
            row.get('Désignation (FR)'), row.get('Designazione (IT)')
        ], {'row': row, 'id': prop_id, 'uri': uri})

    gop_lookup = {}
    for row in gops:
        gop_id = require_id(row, 'GoP-ID', 'GroupOfProperties')
        uri = mint_from_id(base, 'group-of-properties', gop_id)
        add_lookup(gop_lookup, [
            gop_id, row.get('GoP-Code'), row.get('Designation (EN)'), row.get('Beschreibung (DE)'),
            row.get('Description (FR)'), row.get('Descrizione (IT)')
        ], {'row': row, 'id': gop_id, 'uri': uri})

    doc_lookup = {}
    for row in docs:
        doc_id = require_id(row, 'Document-ID', 'Documents')
        uri = mint_from_id(base, 'documents', doc_id)
        add_lookup(doc_lookup, [
            doc_id, row.get('Document-Code'), row.get('DocumentName (EN)'), row.get('Bezeichnung (DE)'),
            row.get('Désignation (FR)'), row.get('Designazione (IT)')
        ], {'row': row, 'id': doc_id, 'uri': uri})

    row2 = {norm(cell_value(wb_data, wb_formula, sheet_name, 2, c)): c for c in range(1, ws.max_column + 1) if norm(cell_value(wb_data, wb_formula, sheet_name, 2, c))}
    property_anchor = row2.get(norm('Property - Designation/Bezeichnung/Désignation/Designazione')) or 5
    document_anchor = row2.get(norm('Document - Designation/Bezeichnung/Désignation/Designazione'))
    loin_anchor = row2.get(norm('LOIN'))
    governance_anchor = row2.get(norm('Governance'))
    property_start_col = property_anchor + 1
    boundaries = [c for c in [document_anchor, loin_anchor, governance_anchor] if c and c > property_start_col]
    property_end_col = min(boundaries) - 1 if boundaries else ws.max_column
    document_start_col = document_anchor + 1 if document_anchor else None
    document_end_col = None
    if document_start_col:
        doc_boundaries = [c for c in [loin_anchor, governance_anchor] if c and c > document_start_col]
        document_end_col = min(doc_boundaries) - 1 if doc_boundaries else ws.max_column

    property_cols = []
    for col in range(property_start_col, property_end_col + 1):
        label = cell_value(wb_data, wb_formula, sheet_name, 2, col)
        if label in (None, ''):
            continue
        prop = prop_lookup.get(norm(label))
        if not prop:
            diagnostics.append(f'{Path(path).name}: WARNING Data_Template row 2 property header does not resolve: {label}')
            continue
        property_cols.append((col, prop, str(label).strip()))

    document_cols = []
    if document_start_col and document_end_col:
        for col in range(document_start_col, document_end_col + 1):
            label = cell_value(wb_data, wb_formula, sheet_name, 2, col)
            if label in (None, ''):
                continue
            doc = doc_lookup.get(norm(label))
            if not doc:
                diagnostics.append(f'{Path(path).name}: WARNING Data_Template row 2 document header does not resolve: {label}')
                continue
            document_cols.append((col, doc, str(label).strip()))

    templates = []
    used_template_ids = set()
    used_requirement_uris = set()
    for row_idx in range(5, ws.max_row + 1):
        row_values = [cell_value(wb_data, wb_formula, sheet_name, row_idx, c) for c in range(1, ws.max_column + 1)]
        if not any(v not in (None, '') for v in row_values):
            continue
        property_assignments = []
        for col, prop, label in property_cols:
            raw = cell_value(wb_data, wb_formula, sheet_name, row_idx, col)
            marker = lexical(raw)
            if not marker:
                continue
            property_assignments.append((col, prop, label, raw, marker))
        selected_docs = []
        for col, doc, label in document_cols:
            marker = lexical(cell_value(wb_data, wb_formula, sheet_name, row_idx, col))
            if marker:
                selected_docs.append(doc)
        group_label = cell_value(wb_data, wb_formula, sheet_name, row_idx, 3)
        group = gop_lookup.get(norm(group_label))
        class_label = cell_value(wb_data, wb_formula, sheet_name, row_idx, 1)
        cls = class_lookup.get(norm(class_label))
        if not property_assignments and not selected_docs and not group:
            continue
        if not cls:
            raise ValueError(f'Data_Template row {row_idx} in {Path(path).name} has assignments but no resolvable Class label')
        if group_label not in (None, '') and not group:
            raise ValueError(f'Data_Template row {row_idx} in {Path(path).name} references unknown GroupOfProperties: {group_label}')
        template_id = f'{row_idx:04d}-{cls["id"]}-{group["id"] if group else "no-group"}'
        if template_id in used_template_ids:
            raise ValueError(f'Data_Template row {row_idx} in {Path(path).name} produces duplicate template ID: {template_id}')
        used_template_ids.add(template_id)
        template_uri = mint_from_id(base, 'data-templates', template_id)
        status = cell_value(wb_data, wb_formula, sheet_name, row_idx, governance_anchor + 1) if governance_anchor else None
        version_date = cell_value(wb_data, wb_formula, sheet_name, row_idx, governance_anchor + 2) if governance_anchor else None
        provenance = cell_value(wb_data, wb_formula, sheet_name, row_idx, governance_anchor + 3) if governance_anchor else None
        loin = {
            'milestone': cell_value(wb_data, wb_formula, sheet_name, row_idx, loin_anchor + 1) if loin_anchor else None,
            'role': cell_value(wb_data, wb_formula, sheet_name, row_idx, loin_anchor + 2) if loin_anchor else None,
            'usecase': cell_value(wb_data, wb_formula, sheet_name, row_idx, loin_anchor + 3) if loin_anchor else None,
        }
        requirements = []
        for col, prop, label, raw, marker in property_assignments:
            prop_id = prop['id']
            req_uri = mint_requirement_uri(base, template_id, prop_id)
            if req_uri in used_requirement_uris:
                raise ValueError(f'Data_Template row {row_idx} in {Path(path).name} creates duplicate requirement for property {prop_id}')
            used_requirement_uris.add(req_uri)
            is_x = marker.casefold() == 'x'
            overrides = [] if is_x else parse_override_list(raw, sheet_name, row_idx, f'col-{col}', diagnostics, path)
            if not is_x:
                override_keys = [lexical(v).casefold() for v in overrides]
                if len(set(override_keys)) != len(override_keys):
                    raise ValueError(f'Data_Template row {row_idx} property {label} has duplicate override values: {overrides}')
            scheme_uri = property_enum_schemes.get(prop_id)
            allowed_value_uris = []
            if overrides:
                if not scheme_uri:
                    raise ValueError(f'Data_Template row {row_idx} property {label} has override {overrides} but no linked Enumeration Scheme')
                lookup = enum_value_lookup.get(prop_id, {})
                for value in overrides:
                    key = lexical(value).casefold()
                    matches = lookup.get(key, [])
                    if len(matches) < 1:
                        raise ValueError(
                            f'Data_Template row {row_idx}; class {cls["id"]}; property {prop_id}; '
                            f'enumeration {scheme_uri}; unresolved override value: {value}'
                        )
                    if len(matches) > 1:
                        diagnostics.append(
                            f'{Path(path).name}: WARNING Data_Template row {row_idx}; class {cls["id"]}; '
                            f'property {prop_id}; override value {value} resolves to {len(matches)} Enumeration Values '
                            'because the source value list has duplicate lexical values.'
                        )
                    allowed_value_uris.extend(matches)
            requirements.append({
                'uri': req_uri,
                'property': prop,
                'marker': marker,
                'uses_complete_scheme': bool(is_x and scheme_uri),
                'scheme_uri': scheme_uri,
                'allowed_value_uris': allowed_value_uris,
            })
        templates.append({
            'id': template_id,
            'uri': template_uri,
            'row': row_idx,
            'class': cls,
            'group': group,
            'documents': selected_docs,
            'requirements': requirements,
            'status': status,
            'version_date': version_date,
            'provenance': provenance,
            'loin': loin,
            'dataset_uri': dataset_uri,
        })
    return templates, diagnostics


def normalize_predefined_type(value):
    if value in (None, ''):
        return ''
    return str(value).strip().replace(' ', '').replace('-', '_').upper()


def ifc_base_class_uri(ifc_entity):
    entity = str(ifc_entity or '').strip()
    if not entity:
        return None
    if not entity.startswith('Ifc'):
        raise ValueError(f'IfcObject Entity must start with Ifc: {entity}')
    return URIRef(f'{IFC_CLASS_BASE}/{entity}')


def ifc_predefined_class_uri(ifc_entity, predefined_type):
    entity = str(ifc_entity or '').strip()
    predefined = normalize_predefined_type(predefined_type)
    if not entity or not predefined or predefined == 'USERDEFINED':
        return None
    if predefined.startswith('Ifc'):
        return URIRef(f'{IFC_CLASS_BASE}/{predefined}')
    return URIRef(f'{IFC_CLASS_BASE}/{entity}{predefined}')


def dataset_uri_for_workbook(path):
    return URIRef(f"https://i14y.admin.ch/catalog/{slug(Path(path).stem)}")


def header_metadata(wb):
    if 'Header' not in wb.sheetnames:
        return {}
    ws = wb['Header']
    meta = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] in (None, ''):
            continue
        key = str(row[0]).strip()
        value = row[1] if len(row) > 1 else None
        if value not in (None, ''):
            meta[key] = value
    return meta


def workbook_resources(wb, wb_path, dict_prefix=''):
    classes = rowdict(wb['Classes']) if 'Classes' in wb.sheetnames else []
    props = rowdict(wb['Properties']) if 'Properties' in wb.sheetnames else []
    gops = rowdict(wb['GroupOfProperties']) if 'GroupOfProperties' in wb.sheetnames else []
    vals = rowdict(wb['Values']) if 'Values' in wb.sheetnames else []
    docs = rowdict(wb['Documents']) if 'Documents' in wb.sheetnames else []
    concept_rel = rowdict(wb['ConceptRelation'], header_row=2, data_start=5) if 'ConceptRelation' in wb.sheetnames else []
    dataset_uri = dataset_uri_for_workbook(wb_path)
    base = LINDAS_BASE if not dict_prefix else f"{LINDAS_BASE}/{dict_prefix}"
    return classes, props, gops, vals, docs, concept_rel, dataset_uri, base


def add_labels(g, uri, de=None, en=None, fr=None, it=None):
    if en: g.add((uri, RDFS.label, Literal(str(en), lang='en')))
    if de: g.add((uri, RDFS.label, Literal(str(de), lang='de')))
    if fr: g.add((uri, RDFS.label, Literal(str(fr), lang='fr')))
    if it: g.add((uri, RDFS.label, Literal(str(it), lang='it')))


def add_resource_links(g, uri, dataset_uri, ifc_uri=None, qudt_uri=None, doc_uri=None):
    g.add((uri, DCTERMS.source, dataset_uri))
    g.add((uri, PROV.wasDerivedFrom, dataset_uri))
    if ifc_uri:
        g.add((uri, SKOS.exactMatch, URIRef(str(ifc_uri))))
    if qudt_uri:
        g.add((uri, QUDT.hasUnit, URIRef(str(qudt_uri))))
    if doc_uri:
        g.add((uri, RDFS.seeAlso, URIRef(str(doc_uri))))


def add_ifc_class_alignment(g, uri, row):
    base_ifc_uri = ifc_base_class_uri(row.get('IfcObject Entity'))
    predefined_ifc_uri = ifc_predefined_class_uri(row.get('IfcObject Entity'), row.get('PredefinedType'))
    if base_ifc_uri is None:
        if row.get('IFC_URI'):
            g.add((uri, SKOS.exactMatch, URIRef(str(row.get('IFC_URI')))))
        return
    g.add((uri, DD.ifcObjectEntity, base_ifc_uri))
    if predefined_ifc_uri is None:
        g.add((uri, SKOS.exactMatch, base_ifc_uri))
        return
    g.add((uri, SKOS.exactMatch, predefined_ifc_uri))
    g.add((uri, SKOS.broadMatch, base_ifc_uri))
    g.add((uri, DD.ifcPredefinedType, predefined_ifc_uri))


def main(argv=None):
    p=argparse.ArgumentParser()
    p.add_argument('kbob')
    p.add_argument('areamgmt')
    p.add_argument('--out', default='out')
    args=p.parse_args(argv)

    kbob=load_workbook(args.kbob,data_only=True)
    area=load_workbook(args.areamgmt,data_only=True)
    kbob_formula=load_workbook(args.kbob,data_only=False)
    area_formula=load_workbook(args.areamgmt,data_only=False)
    outdir=Path(args.out)
    outdir.mkdir(parents=True,exist_ok=True)

    g=Graph()
    g.bind('rdfs', RDFS)
    g.bind('skos', SKOS)
    g.bind('dct', DCTERMS)
    g.bind('prov', PROV)
    g.bind('qudt', QUDT)
    g.bind('owl', OWL)
    g.bind('dd', DD)
    g.bind('schema', SCHEMA)
    g.bind('lindas', Namespace(LINDAS_BASE+'/'))

    diagnostics = []
    # process both workbooks
    for wb, wb_formula, path, prefix in [(kbob, kbob_formula, args.kbob, ''), (area, area_formula, args.areamgmt, BDCH)]:
        classes, props, gops, vals, docs, concept_rel, dataset_uri, base = workbook_resources(wb, path, prefix)
        header = header_metadata(wb)
        enum_schemes, enum_by_designation, enum_diagnostics = build_enumerations(vals)
        diagnostics.extend([f'{Path(path).name}: {line}' for line in enum_diagnostics])
        # dataset provenance node
        g.add((dataset_uri, RDF.type, URIRef('http://www.w3.org/ns/dcat#Dataset')))
        g.add((dataset_uri, RDF.type, DD.DataDictionary))
        dictionary_name_en = header.get('DictionaryName (EN)') or Path(path).stem
        g.add((dataset_uri, RDFS.label, Literal(str(dictionary_name_en), lang='en')))
        if header.get('DictionaryName (DE)'):
            g.add((dataset_uri, RDFS.label, Literal(str(header.get('DictionaryName (DE)')), lang='de')))
        if header.get('DictionaryName (FR)'):
            g.add((dataset_uri, RDFS.label, Literal(str(header.get('DictionaryName (FR)')), lang='fr')))
        if header.get('DictionaryName (IT)'):
            g.add((dataset_uri, RDFS.label, Literal(str(header.get('DictionaryName (IT)')), lang='it')))
        if header.get('DictionaryCode'):
            g.add((dataset_uri, DCTERMS.identifier, Literal(str(header.get('DictionaryCode')))))
        if header.get('DictionaryVersion'):
            g.add((dataset_uri, OWL.versionInfo, Literal(str(header.get('DictionaryVersion')))))
        if header.get('DictionaryUri'):
            g.add((dataset_uri, RDFS.seeAlso, URIRef(str(header.get('DictionaryUri')))))
        g.add((dataset_uri, DD.sourceWorkbook, Literal(Path(path).name)))

        property_uri_by_code = {}
        class_uri_by_code = {}
        gop_uri_by_code = {}
        document_uri_by_name = {}

        for doc in docs:
            doc_id = require_id(doc, 'Document-ID', 'Documents')
            doc_code = str(doc.get('Document-Code') or '').strip()
            doc_uri = mint_from_id(base, 'documents', doc_id)
            for key in [
                doc_id, doc_code, doc.get('DocumentName (EN)'), doc.get('Bezeichnung (DE)'),
                doc.get('Désignation (FR)'), doc.get('Designazione (IT)')
            ]:
                if norm(key):
                    document_uri_by_name[norm(key)] = doc_uri
            g.add((doc_uri, RDF.type, DD.Document))
            g.add((doc_uri, RDF.type, SCHEMA.CreativeWork))
            g.add((dataset_uri, DD.hasDocument, doc_uri))
            add_pref_labels(g, doc_uri, doc.get('Bezeichnung (DE)'), doc.get('DocumentName (EN)'), doc.get('Désignation (FR)'), doc.get('Designazione (IT)'))
            if doc_code:
                g.add((doc_uri, SKOS.notation, Literal(doc_code)))
            for identifier in [doc_id, doc_code]:
                if identifier:
                    g.add((doc_uri, DCTERMS.identifier, Literal(identifier)))
            if doc.get('DocumentOwner'):
                g.add((doc_uri, DCTERMS.publisher, Literal(str(doc.get('DocumentOwner')))))
            if doc.get('Revision'):
                g.add((doc_uri, DD.revision, Literal(str(doc.get('Revision')))))
            if doc.get('Status'):
                g.add((doc_uri, DD.status, Literal(str(doc.get('Status')))))
            if doc.get('Version date'):
                g.add((doc_uri, DCTERMS.modified, Literal(str(doc.get('Version date')))))
            if doc.get('Provenance (PROV)'):
                g.add((doc_uri, PROV.wasAttributedTo, Literal(str(doc.get('Provenance (PROV)')))))
            add_resource_links(g, doc_uri, dataset_uri, doc_uri=doc.get('GUID/URI'))

        for c in classes:
            class_id = require_id(c, 'Class-ID', 'Classes')
            code = str(c.get('Class-Code') or '').strip()
            uri = mint_from_id(base, 'classes', class_id)
            class_uri_by_code[class_id] = uri
            if code:
                class_uri_by_code[code] = uri
            g.add((uri, RDF.type, RDFS.Class))
            g.add((dataset_uri, DD.hasClass, uri))
            add_labels(g, uri, c.get('Bezeichnung (DE)'), c.get('Designation (EN)'), c.get('Désignation (FR)'), c.get('Designazione (IT)'))
            if c.get('Beschreibung (DE)'):
                g.add((uri, SKOS.definition, Literal(str(c.get('Beschreibung (DE)')), lang='de')))
            if c.get('Description (EN)'):
                g.add((uri, SKOS.definition, Literal(str(c.get('Description (EN)')), lang='en')))
            add_resource_links(g, uri, dataset_uri)
            add_ifc_class_alignment(g, uri, c)
            # preserve IFC entity details as labels/comments
            if c.get('IfcObject Entity'):
                g.add((uri, RDFS.comment, Literal(f"IFC entity: {c.get('IfcObject Entity')}")))
            if c.get('PredefinedType'):
                g.add((uri, RDFS.comment, Literal(f"IFC predefined type: {c.get('PredefinedType')}")))
            related_doc = c.get('RelatedDocumentName (EN)') or c.get('RelatedDocuments')
            if related_doc and norm(related_doc) in document_uri_by_name:
                g.add((uri, DD.referencesDocument, document_uri_by_name[norm(related_doc)]))

        for gop in gops:
            gop_id = require_id(gop, 'GoP-ID', 'GroupOfProperties')
            gcode = str(gop.get('GoP-Code') or '').strip()
            gop_uri = mint_from_id(base, 'group-of-properties', gop_id)
            gop_uri_by_code[gop_id] = gop_uri
            if gcode:
                gop_uri_by_code[gcode] = gop_uri
            g.add((gop_uri, RDF.type, DD.GroupOfProperties))
            g.add((dataset_uri, DD.hasGroupOfProperties, gop_uri))
            add_labels(g, gop_uri, gop.get('Beschreibung (DE)'), gop.get('Designation (EN)'), gop.get('Description (FR)'), gop.get('Descrizione (IT)'))
            add_resource_links(g, gop_uri, dataset_uri)

        value_uris_by_scheme_id = {}
        enum_value_lookup_by_scheme_id = {}
        for scheme in enum_schemes:
            scheme_uri = mint_from_id(base, 'enumerations', scheme['id'])
            value_uris = []
            value_lookup = {}
            g.add((scheme_uri, RDF.type, SKOS.ConceptScheme))
            g.add((scheme_uri, RDF.type, DD.EnumerationScheme))
            g.add((dataset_uri, DD.hasEnumerationScheme, scheme_uri))
            add_pref_labels(
                g,
                scheme_uri,
                de=scheme['labels']['de'],
                en=scheme['labels']['en'],
                fr=scheme['labels']['fr'],
                it=scheme['labels']['it'],
            )
            if scheme['code']:
                g.add((scheme_uri, SKOS.notation, Literal(scheme['code'])))
            if scheme['status']:
                g.add((scheme_uri, DD.status, Literal(scheme['status'])))
            add_resource_links(g, scheme_uri, dataset_uri)
            for value in scheme['values']:
                val_uri = mint_value_uri(base, scheme['id'], value['position'], value['canonical'])
                value_uris.append(val_uri)
                g.add((val_uri, RDF.type, SKOS.Concept))
                g.add((val_uri, RDF.type, DD.EnumerationValue))
                g.add((val_uri, RDF.type, DD.AllowedValue))
                g.add((val_uri, SKOS.inScheme, scheme_uri))
                g.add((scheme_uri, DD.hasEnumerationValue, val_uri))
                g.add((val_uri, DD.enumerationPosition, Literal(value['position'], datatype=XSD.integer)))
                add_pref_labels(
                    g,
                    val_uri,
                    de=value['labels'].get('de'),
                    en=value['labels'].get('en'),
                    fr=value['labels'].get('fr'),
                    it=value['labels'].get('it'),
                )
                if value['canonical']:
                    g.add((val_uri, SKOS.notation, Literal(value['canonical'])))
                literal_value = typed_literal(value['raw'])
                if literal_value is not None:
                    g.add((val_uri, RDF.value, literal_value))
                add_resource_links(g, val_uri, dataset_uri)
                for candidate in [value['canonical'], value['raw'], *value['labels'].values()]:
                    key = lexical(candidate).casefold()
                    if key and val_uri not in value_lookup.get(key, []):
                        value_lookup.setdefault(key, []).append(val_uri)
            value_uris_by_scheme_id[scheme['id']] = value_uris
            enum_value_lookup_by_scheme_id[scheme['id']] = value_lookup

        property_enum_schemes = {}
        enum_value_lookup_by_property = {}
        for prop in props:
            prop_id = require_id(prop, 'Property-ID', 'Properties')
            code = str(prop.get('Property-Code') or '').strip()
            uri = mint_from_id(base, 'properties', prop_id)
            property_uri_by_code[prop_id] = uri
            if code:
                property_uri_by_code[code] = uri
            g.add((uri, RDF.type, RDF.Property))
            g.add((dataset_uri, DD.hasProperty, uri))
            add_labels(g, uri, prop.get('Bezeichnung (DE)'), prop.get('Designation (EN)'), prop.get('Désignation (FR)'), prop.get('Designazione (IT)'))
            if prop.get('Beschreibung (DE)'):
                g.add((uri, SKOS.definition, Literal(str(prop.get('Beschreibung (DE)')), lang='de')))
            if prop.get('Description (EN)'):
                g.add((uri, SKOS.definition, Literal(str(prop.get('Description (EN)')), lang='en')))
            add_resource_links(g, uri, dataset_uri, ifc_uri=prop.get('IFC_URI'), qudt_uri=prop.get('QUDT URI'))
            if prop.get('IfcPropertySet (Pset)\nIfcQuantitySet (Qto)'):
                pset_name = str(prop.get('IfcPropertySet (Pset)\nIfcQuantitySet (Qto)'))
                g.add((uri, RDFS.comment, Literal(f"IFC property set: {pset_name}")))
                if pset_name.startswith('http'):
                    pset_uri = URIRef(pset_name)
                    g.add((uri, DD.alignedWithIfcPset, pset_uri))
                else:
                    g.add((uri, DD.alignedWithIfcPset, Literal(pset_name)))
                for gcode, guri in gop_uri_by_code.items():
                    if gcode and (gcode == pset_name or slug(gcode) == slug(pset_name)):
                        g.add((uri, DD.inGroupOfProperties, guri))
                        g.add((guri, DD.groupContainsProperty, uri))
            if prop.get('DataType\n(Base Type)'):
                g.add((uri, DD.datatype, Literal(str(prop.get('DataType\n(Base Type)')))))
                g.add((uri, RDFS.comment, Literal(f"Data type: {prop.get('DataType\n(Base Type)')}")))
            if prop.get('DataType\n(IFC)'):
                g.add((uri, DD.ifcDatatype, Literal(str(prop.get('DataType\n(IFC)')))))
            if prop.get('Status'):
                g.add((uri, DD.status, Literal(str(prop.get('Status')))))
            # preserve enumeration relation
            enum_designation = str(prop.get('EnumerationDesignation (EN)') or '').strip()
            if enum_designation:
                if enum_designation not in enum_by_designation:
                    diagnostics.append(
                        f'{Path(path).name}: WARNING Properties row {prop.get("_row", "?")} references unknown '
                        f'EnumerationDesignation (EN): {enum_designation}'
                    )
                    continue
                scheme = enum_by_designation[enum_designation]
                scheme_uri = mint_from_id(base, 'enumerations', scheme['id'])
                property_enum_schemes[prop_id] = scheme_uri
                enum_value_lookup_by_property[prop_id] = enum_value_lookup_by_scheme_id.get(scheme['id'], {})
                g.add((uri, DD.usesEnumerationScheme, scheme_uri))
                g.add((scheme_uri, DD.enumerationSchemeForProperty, uri))
                for val_uri in value_uris_by_scheme_id.get(scheme['id'], []):
                    g.add((uri, DD.hasAllowedValue, val_uri))
                    g.add((val_uri, DD.valueOfProperty, uri))
                    g.add((uri, SKOS.related, val_uri))

        templates, matrix_diagnostics = build_matrix_model(
            wb, wb_formula, path, base, dataset_uri, classes, props, gops, docs,
            property_enum_schemes, enum_value_lookup_by_property,
        )
        diagnostics.extend(matrix_diagnostics)
        for template in templates:
            t_uri = template['uri']
            cls = template['class']
            group = template['group']
            label_parts = [str(cls['row'].get('Designation (EN)') or cls['row'].get('Bezeichnung (DE)') or cls['id'])]
            if group:
                label_parts.append(str(group['row'].get('Designation (EN)') or group['row'].get('GoP-Code') or group['id']))
            label = ' - '.join(label_parts)
            g.add((t_uri, RDF.type, DD.DataTemplate))
            g.add((dataset_uri, DD.hasDataTemplate, t_uri))
            g.add((t_uri, RDFS.label, Literal(label, lang='en')))
            g.add((t_uri, DCTERMS.identifier, Literal(template['id'])))
            g.add((t_uri, DD.appliesToClass, cls['uri']))
            g.add((cls['uri'], DD.hasDataTemplate, t_uri))
            if group:
                g.add((t_uri, DD.includesGroupOfProperties, group['uri']))
            for doc in template['documents']:
                g.add((t_uri, DD.referencesDocument, doc['uri']))
            if template['status']:
                g.add((t_uri, DD.status, Literal(str(template['status']))))
            if template['version_date']:
                g.add((t_uri, DCTERMS.modified, Literal(str(template['version_date']))))
            if template['provenance']:
                g.add((t_uri, PROV.wasAttributedTo, Literal(str(template['provenance']))))
            for key, value in template['loin'].items():
                if value:
                    g.add((t_uri, getattr(DD, f'loin{key.capitalize()}'), Literal(str(value))))
            add_resource_links(g, t_uri, dataset_uri)
            for requirement in template['requirements']:
                req_uri = requirement['uri']
                prop = requirement['property']
                prop_row = prop['row']
                g.add((req_uri, RDF.type, DD.PropertyRequirement))
                g.add((t_uri, DD.hasPropertyRequirement, req_uri))
                g.add((req_uri, DD.requirementOfDataTemplate, t_uri))
                g.add((req_uri, DD.appliesToClass, cls['uri']))
                g.add((req_uri, DD.requiresProperty, prop['uri']))
                g.add((prop['uri'], DD.requiredByTemplate, req_uri))
                if group:
                    g.add((req_uri, DD.inGroupOfProperties, group['uri']))
                for doc in template['documents']:
                    g.add((req_uri, DD.referencesDocument, doc['uri']))
                g.add((req_uri, DD.inclusionMarker, Literal(str(requirement['marker']))))
                g.add((req_uri, DD.requirementLevel, Literal('included')))
                if prop_row.get('DataType\n(Base Type)'):
                    g.add((req_uri, DD.contextualDatatype, Literal(str(prop_row.get('DataType\n(Base Type)')))))
                if prop_row.get('QUDT URI'):
                    g.add((req_uri, QUDT.hasUnit, URIRef(str(prop_row.get('QUDT URI')))))
                    g.add((req_uri, DD.contextualUnit, URIRef(str(prop_row.get('QUDT URI')))))
                elif prop_row.get('Units/Einheit/Unité/Unità'):
                    g.add((req_uri, DD.contextualUnit, Literal(str(prop_row.get('Units/Einheit/Unité/Unità')))))
                if requirement['scheme_uri']:
                    g.add((req_uri, DD.usesEnumerationScheme, requirement['scheme_uri']))
                    g.add((req_uri, DD.usesCompleteValueScheme, Literal(requirement['uses_complete_scheme'])))
                for val_uri in requirement['allowed_value_uris']:
                    g.add((req_uri, DD.allowedValue, val_uri))
                add_resource_links(g, req_uri, dataset_uri)

        # preserve concept relations exactly as declared
        for row in concept_rel:
            s_code = str(row.get('PropertyCode') or '').strip()
            rel = str(row.get('RelationType') or '').strip()
            target = str(row.get('RelatedPropertyUri') or '').strip()
            if not s_code or not rel or not target:
                continue
            if s_code not in property_uri_by_code and s_code not in class_uri_by_code:
                raise ValueError(f'ConceptRelation row {row.get("_row", "?")} references unknown source code/id: {s_code}')
            s_uri = property_uri_by_code.get(s_code) or class_uri_by_code.get(s_code)
            pred_local = rel.split(':')[-1]
            if hasattr(SKOS, pred_local):
                pred = getattr(SKOS, pred_local)
            else:
                pred = URIRef(rel)
            t_uri = URIRef(target) if target.startswith('http') else mint_from_id(base, 'properties', target)
            g.add((s_uri, pred, t_uri))
            if row.get('Notes'):
                g.add((s_uri, RDFS.comment, Literal(str(row.get('Notes')))))

    out_ttl=outdir/'combined.ttl'
    g.serialize(destination=str(out_ttl), format='turtle')
    diagnostics_file = outdir/'mapper-diagnostics.md'
    diagnostics_file.write_text(
        '# Mapper Diagnostics\n\n'
        '## Enumeration publication\n\n'
        + ('\n'.join(f'- {line}' for line in diagnostics) if diagnostics else '- No enumeration schemes found.')
        + '\n',
        encoding='utf-8',
    )
    print('wrote', out_ttl)
    print('wrote', diagnostics_file)

if __name__=='__main__':
    import sys
    main(sys.argv[1:])
